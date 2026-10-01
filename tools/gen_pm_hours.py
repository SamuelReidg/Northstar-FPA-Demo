"""Write the synthetic PM Hours & Expenses export ("CM PM Report.xlsx") from the demo world.

    python3 tools/gen_pm_hours.py

Output: demo_data/Weekly and Monthly Updates/9.29.2026/PM Hours & Expenses/CM PM Report.xlsx

Layout mirrors the ERP export that pipeline/regenerate_pm_hours_data.py reads:
  * 'Timesheet Data' (23 columns) - one row per daily time record in W.entries (FY2026
    through the as-of week, including the Dec-29..31 days of the first week), joined to the
    employee (home location, labor cost type, approver) and project (billing type, PM, name).
    Overhead rows carry Project ID "OVHD" and the pseudo-project name "Overhead Global".
  * 'Expense Data' (22 columns) - one row per expense line on a POSTED report (Approved or
    Paid) whose GL posting date falls on or before the as-of date. Transaction date = the
    line's date, GL Posting Date = the report's GL date; project / PM / customer come from
    the project the line was charged to, and the task is the one the employee logged the most
    hours on for that project that week.

Scope: "CM" = the Central & Mountain region. Like the production report, the export covers only
that region's project managers (CM_OFFICES below: about half the company's project hours), which
also keeps the dashboard's data file small enough to open quickly. Timesheet rows are kept when
the project's PM belongs to a CM office, and overhead rows when the employee's home office does;
expense lines follow the project's PM.

Everything is derived from tools/demo_world.py - nothing is random except the free-text
flavour (notes, vendors), which uses a seeded RNG so the file is deterministic.
"""
from __future__ import annotations

import datetime as dt
import os
import random
import re
import sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from demo_world import world  # noqa: E402

from openpyxl import Workbook  # noqa: E402

OUT_DIR = os.path.join(ROOT, 'demo_data', 'Weekly and Monthly Updates', '9.29.2026', 'PM Hours & Expenses')
OUT = os.path.join(OUT_DIR, 'CM PM Report.xlsx')

TS_HEADER = ['Legacy ID', 'Contact name', 'Location name', 'Entry date', 'Project ID', 'Utilized hours',
             'Non-utilized hours', 'Hours', 'Utilization Rate', 'Time type', 'Task name', 'Notes',
             'Billing type', 'Cost type ID', 'User ID', 'Username', 'Cost type ID', 'Item ID', 'Item name',
             'Name', 'Standard cost type name', 'Project Manager Name', 'Project name']
EX_HEADER = ['Transaction Date', 'GL Posting Date', 'Location ID', 'Location Name', 'Employee ID',
             'Employee Name', 'Amount', 'ER #', 'ER Title', 'Ln #', 'GL Account #', 'GL Account Name',
             'Paid To', 'Paid For', 'Customer ID', 'Customer Name', 'Project Manager ID',
             'Project Manager Name', 'Project ID', 'Project Name', 'Task ID', 'Task Name']

OVERHEAD_PROJECT_NAME = 'Overhead Global'

# Central & Mountain region: the PMOs this report covers (~50% of project hours).
CM_OFFICES = {'Big Sky', 'Front Range', 'High Desert', 'Red Rock', 'Prairie', 'Gulf Coast', 'Hill Country',
              'Built Environment', 'Paleontology', 'Meridian Central', 'Meridian Southwest'}

# Overhead task -> (cost type ID, item ID, item name). Anything not listed is administrative.
OVHD_COST = [
    ('Meetings',        ('OH-MTG', 'OH-MTGS', 'Meetings')),
    ('Marketing',       ('OH-MKT', 'OH-MKTG', 'Marketing')),
    ('Recruiting',      ('OH-RCT', 'OH-RCRT', 'Recruiting')),
    ('Expense Reports', ('OH-EXR', 'OH-EXRPT', 'Expense Reports')),
    ('Required Training', ('OH-TRN', 'OH-TRNG', 'Training')),
    ('Vacation',        ('OH-TOFF', 'OH-TOFF', 'Time Off')),
    ('PTO',             ('OH-TOFF', 'OH-TOFF', 'Time Off')),
    ('Holiday',         ('OH-HOL', 'OH-HOL', 'Holiday')),
    ('Sick',            ('OH-SCK', 'OH-SCK', 'Sick')),
    ('Parental Leave',  ('OH-LV', 'OH-LV', 'Leave')),
]
OVHD_ADMIN = ('OH-ADM', 'OH-ADM', 'Administrative')

# Expense type (world) -> GL account (number, name). Names follow a generic services chart of accounts.
GL_OF = {
    'Per Diem':         (51430, 'Per Diem - Billable'),
    'Lodging':          (51200, 'Hotel & Lodging - Billable'),
    'Mileage':          (51600, 'Mileage Reimbursement - Billable'),
    'Field Supplies':   (50100, 'Field Materials & Supplies - Billable'),
    'Vehicle Rental':   (51000, 'Travel & Transportation - Billable'),
    'Airfare':          (51000, 'Travel & Transportation - Billable'),
    'Equipment Rental': (50130, 'Equipment Rental - Billable'),
}
VENDORS = {
    'Per Diem':         [None],
    'Lodging':          ['Trailhead Inn', 'Riverside Motel', 'Canyon Lodge', 'Mesa Suites', 'Pine Hollow Motor Inn'],
    'Mileage':          [None],
    'Field Supplies':   ['Hardware store', 'Outdoor supply', 'Survey supply catalog', 'Farm & ranch supply'],
    'Vehicle Rental':   ['Rental car agency', 'Truck rental', 'Fuel station'],
    'Airfare':          ['Airline'],
    'Equipment Rental': ['Equipment rental yard', 'Tool rental'],
}

NOTES = {
    'Field': ['Field day', 'field day', 'field work', 'Field survey'],
    'Survey': ['Survey', 'survey', 'pedestrian survey'],
    'Monitoring': ['Monitoring', 'monitoring', 'construction monitoring'],
    'Report': ['Reporting', 'reporting', 'report edits', 'draft report'],
    'Research': ['background research', 'records search', 'research'],
    'GIS': ['GIS', 'mapping', 'figures'],
    'Lab': ['lab', 'artifact analysis', 'catalog'],
    'Project Management': ['PM', 'project management', 'billing review', 'client call'],
    'Vacation': ['PTO', 'Vacation'],
    'PTO': ['PTO'],
    'Holiday': ['Holiday'],
    'Sick': ['Sick'],
    'Training': ['training', 'required training'],
    'Meetings': ['staff meeting', 'meetings', 'team meeting'],
    'Proposal': ['proposal', 'proposal writing', 'cost estimate'],
    'Administrative': ['admin', 'Admin', 'email', '.'],
}


def num(x):
    """Whole numbers as int (the export writes 8, not 8.0); others as float."""
    return int(x) if float(x).is_integer() else float(x)


def note_for(task, rnd):
    for key, opts in NOTES.items():
        if key.lower() in task.lower():
            return rnd.choice(opts)
    return rnd.choice(['.', 'admin', 'misc', 'coordination'])


def labor_category(e):
    """Standard cost type (labor category) for an employee, from title/role."""
    if e['role'] == 'admin':
        return 'Administrative Assistant/Clerical'
    if e['role'] == 'lead':
        return 'Principal Investigator/Lead'
    t = re.sub(r'\s+(I|II|III)$', '', e['title'])
    if t.startswith('Staff '):
        return 'Staff Level'
    if t in ('Team Lead', 'Regional Principal Investigator'):
        return 'Principal Investigator/Lead'
    if t.startswith('International'):
        return 'Senior Level'
    return t


def first_last(name):
    last, first = [x.strip() for x in name.split(',', 1)]
    return '%s %s' % (first, last)


def main():
    W = world()
    rnd = random.Random('%s-pm_hours' % W.seed)
    by_name = {e['name']: e for e in W.employees}

    cats = sorted({labor_category(e) for e in W.employees})
    cat_code = {c: 14010 + 7 * i for i, c in enumerate(cats)}

    def approver(e):
        sup = by_name.get(e['supervisor'] or '') or next(x for x in W.employees if x['office'] == 'Corporate')
        return sup['email'].split('@')[0], first_last(sup['name'])

    def ovhd_cost(task):
        for key, ct in OVHD_COST:
            if key.lower() in task.lower():
                return ct
        return OVHD_ADMIN

    wb = Workbook(write_only=True)
    ts = wb.create_sheet('Timesheet Data')
    ts.append(TS_HEADER)
    n_ts = 0
    # task most used per (emp, week, project): feeds the expense lines' task
    task_hours = defaultdict(lambda: defaultdict(float))
    def cm_project(pid):
        return by_name[W.project[pid]['pm']]['office'] in CM_OFFICES

    for r in W.entries:
        e = W.emp[r['emp']]
        if (e['office'] not in CM_OFFICES) if r['project'] == 'OVHD' else not cm_project(r['project']):
            continue
        loc = W.office[e['office']]['location']
        uid, uname = approver(e)
        if r['project'] == 'OVHD':
            ct, item_id, item = ovhd_cost(r['task'])
            billing, pm, pname = None, None, OVERHEAD_PROJECT_NAME
        else:
            p = W.project[r['project']]
            item = labor_category(e)
            ct = item_id = cat_code[item]
            billing, pm, pname = p['billing'], first_last(p['pm']), p['name']
            task_hours[(r['emp'], r['week'], r['project'])][r['task']] += r['hours']
        util = r['util']
        eid = e['id'] if isinstance(e['id'], int) else str(e['id'])
        ts.append([eid, e['name'], loc, dt.datetime.combine(r['date'], dt.time()), r['project'],
                   num(util), num(round(r['hours'] - util, 2)), num(r['hours']),
                   100 if util else 0, r['tt'], r['task'], note_for(r['task'], rnd), billing,
                   ct, uid, uname, ct, item_id, item, item, item, pm, pname])
        n_ts += 1

    # ── Expense Data: posted (Approved/Paid) report lines, GL date <= as-of ──
    reports = {x['er']: x for x in W.expense_reports}
    client_id = {c: 11000 + 3 * i for i, c in enumerate(W.clients)}
    task_id = {}
    nxt = [40000]

    def tid(pid, task):
        k = (pid, task)
        if k not in task_id:
            nxt[0] += 1 + (len(task_id) * 7) % 5
            task_id[k] = 'PCT-%06d' % nxt[0]
        return task_id[k]

    ex = wb.create_sheet('Expense Data')
    ex.append(EX_HEADER)
    ln = defaultdict(int)
    n_ex, total = 0, 0.0
    for l in W.expense_lines:
        rep = reports[l['er']]
        if rep['state'] not in ('Approved', 'Paid') or rep['gl'] > W.as_of:
            continue
        if not cm_project(l['project']):
            continue
        e = W.emp[l['emp']]
        o = W.office[e['office']]
        p = W.project[l['project']]
        pm = by_name[p['pm']]
        gl_no, gl_name = GL_OF[l['type']]
        ln[l['er']] += 1
        th = task_hours.get((l['emp'], rep['week'], l['project']))
        task = max(th, key=th.get) if th else p['tasks'][0]
        t_id = tid(p['id'], task)
        title = rnd.choice(['%s %s' % (p['id'], l['type']),
                            '%s %s' % (rep['reason'], rep['week'].strftime('%m/%d')),
                            '%s - %s' % (rep['reason'], p['name'])])
        vendor = rnd.choice(VENDORS[l['type']])
        paid_for = rnd.choice(['Field crew', 'Site visit', None, None, None]) if vendor else None
        ex.append([dt.datetime.combine(l['date'], dt.time()), dt.datetime.combine(rep['gl'], dt.time()),
                   o['location_id'], o['location'], e['sage_id'], e['name'], round(l['amount'], 2),
                   l['er'], title, ln[l['er']], gl_no, gl_name, vendor, paid_for,
                   client_id[p['client']], p['client'], pm['sage_id'], first_last(pm['name']),
                   p['id'], p['name'], t_id, '%s - %s' % (task, t_id)])
        n_ex += 1
        total += l['amount']

    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(OUT)
    print('Wrote %s' % OUT)
    print('  Timesheet Data rows: %d' % n_ts)
    print('  Expense Data rows:   %d  (posted total %.2f)' % (n_ex, total))


if __name__ == '__main__':
    main()
