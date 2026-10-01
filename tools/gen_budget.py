#!/usr/bin/env python3
"""gen_budget.py -- synthetic "FP&A & QC Project Report" exports for the Budget Watchtower.

Writes, from the shared synthetic world (tools/demo_world.py), in exactly the layout the
pipeline's loader (pipeline/regenerate_budget_tabs.py) expects:

  demo_data/Weekly and Monthly Updates/9.29.2026/Budget Watchtower/
      FP&A & QC Project Report (41).xlsx        this week, USD, "As of September 29, 2026"
      FP&A & QC Project Report E3000.xlsx       companion (Gulf entity), SAR, same "As of" date
  demo_data/Weekly and Monthly Updates/9.22.2026/Budget Watchtower/
      FP&A & QC Project Report (40).xlsx        last week, USD, "As of September 22, 2026"
  demo_data/fixtures/budget_regression.xlsx     archived pre-08/2026-layout workbook
                                                ("All Active Projects", "As of June 16, 2026")
                                                used by the loader regression test (--check)

Every figure is derived from the world, never drawn independently:
  * cost to date   = prior-year labor + FY labor (timesheet hours x cost rate, OT x1.5),
                     grossed up for the project's expense share -- identical to the world's
                     cost_to_date on the current as-of;
  * earlier as-ofs = the same roll-up with the later weeks' W.entries backed out (last week's
                     report drops the week ending 2026-09-26; the June fixture drops everything
                     after the week ending 2026-06-13);
  * revenue        = fixed fee: contract x min(1, cost / revised budget); T&M: prior revenue +
                     billed hours x bill rate (capped at contract) -- the world's own rule;
  * invoiced       = revenue less the world's unbilled gap for that project;
  * Budget Spent % = cost / revised budget (4 dp), REV REC % = revenue / contract (3 dp),
    Gross profit   = contract - cost, Margin % = gross profit / contract (the export's own maths);
  * OVHD rows and admin cost buckets carry the overhead hours actually booked to project "OVHD".

The E3000 companion: International "Transferred Entity" projects keep their budgets and their
pre-transfer cost in the USD report, but contract, revenue, invoicing and post-transfer
(September) cost are $0 there and appear instead in the E3000 file, in SAR (x3.75), with IDs
suffixed "-Int". One transferred project is listed under both "ID" and "ID-Int"; one project
exists only in E3000 (also split "ID" / "ID-Int"); two transferred projects are missing from
E3000 (held by another entity) and so drop out on the commercial-basis rule.

Usage:
    python tools/gen_budget.py            # write all four workbooks
    python tools/gen_budget.py --check    # recompute the fixture through the pipeline and
                                          # assert the documented regression numbers
"""
import datetime as dt
import os
import random
import sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from demo_world import world, NON_POSTING, COMPANY   # noqa: E402

from openpyxl import Workbook                              # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment   # noqa: E402

PKG = os.path.join(ROOT, 'demo_data', 'Weekly and Monthly Updates')
CUR_DIR = os.path.join(PKG, '9.29.2026', 'Budget Watchtower')
PREV_DIR = os.path.join(PKG, '9.22.2026', 'Budget Watchtower')
FIX_DIR = os.path.join(ROOT, 'demo_data', 'fixtures')
CUR_FILE = os.path.join(CUR_DIR, 'FP&A & QC Project Report (41).xlsx')
E3000_FILE = os.path.join(CUR_DIR, 'FP&A & QC Project Report E3000.xlsx')
PREV_FILE = os.path.join(PREV_DIR, 'FP&A & QC Project Report (40).xlsx')
FIX_FILE = os.path.join(FIX_DIR, 'budget_regression.xlsx')

SAR_PER_USD = 3.75
TRANSFER_DATE = dt.date(2026, 9, 1)     # Gulf entity takes over from September

# Report dates (the export's "As of") and the last week-ending Saturday each one covers.
CUR = dict(report=dt.date(2026, 9, 29), cutoff=dt.date(2026, 9, 26))
PREV = dict(report=dt.date(2026, 9, 22), cutoff=dt.date(2026, 9, 19))
FIX = dict(report=dt.date(2026, 6, 16), cutoff=dt.date(2026, 6, 13))

ACTIVE = {'In Progress', 'Transferred Entity', 'Final Reporting/Review', 'Lab/Curation',
          'Cost Tracking on Closed Project'}
CLOSED = ('Completed', 'Final Billing/Invoice', 'Cost Tracking on Closed Project')
BUCKET_TASKS = {   # overhead tasks booked to admin cost buckets instead of the office OVHD row
    'OVHD - Labor - Proposal Writing': 'MKTG', 'OVHD - Production - Proposal Editing': 'MKTG',
    'OVHD - Labor - Business Development': 'MKTG', 'OVHD - Labor - Marketing': 'MKTG',
    'OVHD - Labor - Company Meeting': 'MEET',
}

# ----------------------------------------------------------------------------- styling
FONT = 'Helvetica'
F_BODY = Font(name=FONT, size=10)
F_TITLE = Font(name=FONT, size=14)
WHITE = PatternFill('solid', fgColor='FFFFFFFF')
RED = PatternFill('solid', fgColor='FFF7A3A5')
MONEY = '#,##0.00;(#,##0.00)'
PCT1 = '#,##0.0 %;(#,##0.0) %'
PCT2 = '#,##0.00 %;(#,##0.00) %'
LEFT = Alignment(horizontal='left')
RIGHT = Alignment(horizontal='right')
WIDTHS = [46, 110.4, 28, 21, 23.4, 23, 15.6, 15.7, 13, 13, 23.9, 14, 21.7, 15.7, 13, 13.3,
          14.9, 14.1, 13.9, 13, 12.7, 14.9]


def fmt_date(d):
    return d.strftime('%m/%d/%Y')


def long_date(d):
    return '%s %d, %d' % (d.strftime('%B'), d.day, d.year)


def month_end(d):
    nxt = d.replace(day=28) + dt.timedelta(days=4)
    return nxt - dt.timedelta(days=nxt.day)


def display_name(last_first):
    if ', ' in last_first:
        last, first = last_first.split(', ', 1)
        return '%s %s' % (first, last)
    return last_first


# ----------------------------------------------------------------------------- world roll-up
class Ledger:
    """Per-project labor / billing by week (and month), from W.entries."""

    def __init__(self, W):
        self.W = W
        self.lab = defaultdict(float)        # (pid, week) -> labor $
        self.lab_m = defaultdict(float)      # (pid, week, (y, m)) -> labor $
        self.lab_pre = defaultdict(float)    # (pid, week) -> labor $ dated before TRANSFER_DATE
        self.bill = defaultdict(float)       # (pid, week) -> billed $
        self.fy_total = defaultdict(float)   # pid -> full-year labor
        self.ovhd = defaultdict(float)       # (bucket key, week) -> labor $
        self.ovhd_m = defaultdict(float)     # (bucket key, week, (y, m))
        self.ovhd_pre = defaultdict(float)
        for r in W.entries:
            e = W.emp[r['emp']]
            mult = 1.5 if r['tt'] == 'OT' else 1.0
            amt = r['hours'] * e['cost_rate'] * mult
            ym = (r['date'].year, r['date'].month)
            if r['project'] == 'OVHD':
                key = BUCKET_TASKS.get(r['task']) or ('OVHD', e['office'])
                self.ovhd[(key, r['week'])] += amt
                self.ovhd_m[(key, r['week'], ym)] += amt
                if r['date'] < TRANSFER_DATE:
                    self.ovhd_pre[(key, r['week'])] += amt
                continue
            pid = r['project']
            self.lab[(pid, r['week'])] += amt
            self.lab_m[(pid, r['week'], ym)] += amt
            if r['date'] < TRANSFER_DATE:
                self.lab_pre[(pid, r['week'])] += amt
            self.fy_total[pid] += amt
            if r['task'] != NON_POSTING:
                self.bill[(pid, r['week'])] += r['hours'] * e['bill_rate']
        self.weeks = sorted({w for (_, w) in self.lab} | {w for (_, w) in self.ovhd})

    def upto(self, table, key, cutoff):
        return sum(table.get((key, w), 0.0) for w in self.weeks if w <= cutoff)

    def month(self, table, key, cutoff, ym):
        return sum(table.get((key, w, ym), 0.0) for w in self.weeks if w <= cutoff)

    def later(self, pid, cutoff):
        return any(self.lab.get((pid, w), 0.0) for w in self.weeks if w > cutoff)


def status_at(p, L, snap):
    """Project status as the export would have shown it on that as-of."""
    st = p['status']
    if snap is FIX:
        # Before the 09/2026 vocabulary change there was no "Transferred Entity"; jobs that
        # closed out after June were still running then.
        if st == 'Transferred Entity':
            return 'In Progress'
        if st in ('Completed', 'Final Billing/Invoice', 'On Hold') and L.later(p['id'], snap['cutoff']):
            return 'In Progress'
    return st


def figures(p, L, snap, status):
    """Contract, revenue, invoiced, budgets and cost of one project on one as-of (USD)."""
    W = L.W
    pid, rb, es = p['id'], p['revised_budget'], p['expense_share']
    c = snap['cutoff']
    fy_lab = L.upto(L.lab, pid, c)
    prior = p['prior_share'] * rb * (1 - es) * (1 + p['burn_bias'])   # world's prior-year labor
    if p['status'] in CLOSED and L.fy_total[pid] == 0:
        prior = rb * (1 - es) * min(1.3, max(0.7, 1 + p['burn_bias']))
    gross = 1.0 / (1 - es)
    cost = (prior + fy_lab) * gross
    billed = L.upto(L.bill, pid, c)

    def revenue_for(cst, bld, st):
        if p['billing'] == 'Fixed fee':
            pct = min(1.0, round(cst / rb, 4)) if rb else 0.0   # world rounds % complete to 4 dp
            if st in ('Completed', 'Final Billing/Invoice'):
                pct = 1.0
            return p['contract'] * pct
        return min(p['contract'], p['prior_share'] * p['contract'] + bld)

    rev = revenue_for(cost, billed, status)
    gap = p['revenue'] - p['invoiced']   # the world's unbilled lag (negative = billed in excess)
    inv = max(0.0, rev - gap)
    rev0 = revenue_for(prior * gross, 0.0, 'In Progress')  # at the start of the year
    ym = (snap['report'].year, snap['report'].month)
    return dict(
        contract=p['contract'], revenue=rev, invoiced=inv,
        revbud=rb, orig=p['orig_budget'], etc=rb - p['orig_budget'], cost=cost,
        ytd=fy_lab * gross, month=L.month(L.lab_m, pid, c, ym) * gross,
        inv_ytd=max(0.0, inv - rev0),
        pre_cost=(prior + L.upto(L.lab_pre, pid, c)) * gross,   # cost booked before the transfer
    )


# ----------------------------------------------------------------------------- choices
def rng_for(W, tag):
    return random.Random('%s-budget-%s' % (W.seed, tag))


def location_for(W, p):
    """Export 'Location Name': mostly the office's ERP location, with the sub-locations the
    dashboard's MAP rolls back up into their PMO (see build_overrun_dashboard.MAP)."""
    r = rng_for(W, 'loc-' + p['id']).random()
    o, yr = p['office'], int(p['id'][:2])
    if o == 'Sierra':
        return 'Sierra North' if r < 0.35 else 'Sierra South' if r < 0.65 else 'Sierra'
    if o == 'Gulf Coast':
        return 'Gulf Coast East' if r < 0.4 else 'Gulf Coast West' if r < 0.75 else 'Gulf Coast'
    if o == 'Piedmont' and yr <= 24:     # retired office code, still on older jobs
        return 'PIEDMONT - INACTIVE' if r < 0.5 else 'Piedmont - Inactive'
    if o == 'Meridian Central' and yr <= 24:   # pre-integration location of the acquired firm
        return 'Meridian'
    return p['location']


def ontrack(W, p, snap, bs, status):
    """'On-Track As of Date': the last PM on-track review, recorded on some late-stage jobs."""
    if status not in ACTIVE or bs < 0.7:
        return ' '
    r = rng_for(W, 'ot-' + p['id'])
    if r.random() > 0.55:
        return ' '
    fridays = [dt.date(2026, 7, 31) + dt.timedelta(days=7 * i) for i in range(8)]
    pick = fridays[int(r.random() * len(fridays))]
    while pick > snap['cutoff']:
        pick -= dt.timedelta(days=14)
    return fmt_date(pick)


def intl_roles(W, L):
    """Which International projects demo which E3000 merge case (deterministic)."""
    intl = sorted((p for p in W.projects if p['office'] == 'International'), key=lambda p: p['id'])
    te = [p for p in intl if p['status'] == 'Transferred Entity']
    sept = lambda p: L.month(L.lab_m, p['id'], CUR['cutoff'], (2026, 9))
    missing = [p['id'] for p in te if L.fy_total[p['id']] == 0 and p['cost_to_date'] > 0][:2]
    dual = max((p for p in te if p['id'] not in missing), key=sept)['id']
    only = max((p for p in intl if p['status'] == 'In Progress'), key=sept)['id']
    return dict(missing=set(missing), dual=dual, only=only,
                te=[p['id'] for p in te if p['id'] not in missing])


# ----------------------------------------------------------------------------- row building
def row(pid, name, status, billing, location, pm, ot, f):
    return dict(id=pid, name=name, status=status, billing=billing, location=location, pm=pm,
                ot=ot, H=f['contract'], I=f['revenue'], J=f['invoiced'], K=f['revbud'],
                L=f['orig'], M=f['etc'], N=f['cost'], T=f['ytd'], U=f['month'], V=f['inv_ytd'])


ZERO = dict(contract=0.0, revenue=0.0, invoiced=0.0, revbud=0.0, orig=0.0, etc=0.0, cost=0.0,
            ytd=0.0, month=0.0, inv_ytd=0.0)


def office_lead(W, office):
    ppl = [e for e in W.employees if e['office'] == office and e['status'] == 'Active']
    lead = next((e for e in ppl if e['role'] == 'lead'), None) or (ppl[0] if ppl else None)
    return display_name(lead['name']) if lead else ' '


def overhead_rows(W, L, snap, intl_split):
    """Office OVHD rows + admin cost buckets (no contract, no budget)."""
    out = []
    ym = (snap['report'].year, snap['report'].month)
    for o in W.offices:
        key = ('OVHD', o['name'])
        cost = L.upto(L.ovhd, key, snap['cutoff'])
        if not cost:
            continue
        month = L.month(L.ovhd_m, key, snap['cutoff'], ym)
        if intl_split and o['name'] == 'International':
            cost = L.upto(L.ovhd_pre, key, snap['cutoff'])
            month = 0.0
        f = dict(ZERO, cost=cost, ytd=cost, month=month)
        out.append(row('%s OVHD' % o['name'], '%s Overhead' % o['name'], 'In Progress', ' ',
                       o['location'], office_lead(W, o['name']), ' ', f))
    sm = next(e for e in W.employees if e['office'] == 'Sales & Marketing' and e['status'] == 'Active')
    corp = next(e for e in W.employees if e['office'] == 'Corporate' and e['status'] == 'Active')
    for key, pid, name, loc, pm in (('MKTG', 'MKTG-2026', 'Sales and Marketing', 'Sales and Marketing', sm),
                                    ('MEET', 'EVT-2026-01', 'Company Meeting 2026', 'Corporate', corp)):
        cost = L.upto(L.ovhd, key, snap['cutoff'])
        f = dict(ZERO, cost=cost, ytd=cost, month=L.month(L.ovhd_m, key, snap['cutoff'], ym))
        out.append(row(pid, name, 'In Progress', 'Time and material', loc, display_name(pm['name']), ' ', f))
    return out


def pending_rows(W, snap):
    """'Active Projects' group of the export: awards set up but not yet billing-active."""
    wins = sorted((o for o in W.opportunities if o['stage'] == 'Closed Won'
                   and snap['report'] - dt.timedelta(days=14) <= o['close'] <= snap['report']),
                  key=lambda o: o['id'])
    base = max(int(p['id'][-5:]) for p in W.projects)
    out = []
    for i, o in enumerate(wins):
        name = o['name'].split(' | ', 1)[-1]
        office = next(x for x in W.offices if x['name'] == o['office'])
        f = dict(ZERO, contract=o['amount'])
        out.append(row('26-PC-%05d' % (base + 7 + i * 3), name, 'Pending Billing Activation',
                       'Time and material' if int(o['id'][-1]) % 2 else 'Fixed fee',
                       office['location'], display_name(o['owner']), ' ', f))
    return out


def build_rows(W, L, snap, with_companion):
    roles = intl_roles(W, L)
    main, comp = [], []
    for p in sorted(W.projects, key=lambda p: p['id']):
        if p['start'] > snap['report']:
            continue
        status = status_at(p, L, snap)
        f = figures(p, L, snap, status)
        bs = f['cost'] / f['revbud'] if f['revbud'] else 0.0
        loc = location_for(W, p)
        pm = display_name(p['pm'])
        ot = ontrack(W, p, snap, bs, status)
        if p['office'] == 'International':
            loc = 'Global LLC International' if p['id'] in roles['missing'] else loc
            if with_companion and p['id'] == roles['only']:
                # Created in the Gulf entity: E3000 only, split over "ID" and "ID-Int".
                post = f['cost'] - f['pre_cost']
                a = dict(f, cost=f['pre_cost'], ytd=f['ytd'] - post, month=0.0)
                b = dict(ZERO, cost=post, ytd=post, month=post)
                comp.append(row(p['id'], p['name'], 'In Progress', p['billing'], 'NH Gulf', pm, ' ', a))
                comp.append(row(p['id'] + '-Int', p['name'], 'In Progress', p['billing'], 'NH Gulf', pm, ' ', b))
                continue
            if with_companion and p['status'] == 'Transferred Entity':
                # USD report keeps the budget and pre-transfer cost; everything else moves.
                post = f['cost'] - f['pre_cost']
                keep = dict(ZERO, revbud=f['revbud'], orig=f['orig'], etc=f['etc'],
                            cost=f['pre_cost'], ytd=f['ytd'] - post)
                main.append(row(p['id'], p['name'], status, p['billing'], loc, pm, ot, keep))
                if p['id'] in roles['missing']:
                    continue
                moved = dict(f, cost=post, ytd=post, month=post)
                if p['id'] == roles['dual']:
                    # listed twice in E3000: commercial terms under "ID-Int", September cost under "ID"
                    comp.append(row(p['id'] + '-Int', p['name'], 'In Progress', p['billing'], 'NH Gulf', pm, ' ',
                                    dict(moved, cost=0.0, ytd=0.0, month=0.0)))
                    comp.append(row(p['id'], p['name'], 'In Progress', p['billing'], 'NH Gulf', pm, ' ',
                                    dict(ZERO, cost=post, ytd=post, month=post)))
                else:
                    comp.append(row(p['id'] + '-Int', p['name'], 'In Progress', p['billing'], 'NH Gulf', pm, ' ', moved))
                continue
        main.append(row(p['id'], p['name'], status, p['billing'], loc, pm, ot, f))
    if with_companion:
        post = (L.upto(L.ovhd, ('OVHD', 'International'), snap['cutoff'])
                - L.upto(L.ovhd_pre, ('OVHD', 'International'), snap['cutoff']))
        comp.append(row('OVHD', ' ', 'In Progress', ' ', ' ', ' ', ' ',
                        dict(ZERO, cost=post, ytd=post, month=post)))
    main += overhead_rows(W, L, snap, with_companion)
    main.sort(key=lambda r: (r['id'][0].isalpha(), r['id']))
    return main, comp, roles


# ----------------------------------------------------------------------------- writers
def _derived(r):
    H, I, K, N = r['H'], r['I'], r['K'], r['N']
    R = H - N
    return dict(O=K - N, P=round(I / H, 3) if H else 0, Q=round(N / K, 4) if K else 0,
                R=R, S=round(R / H, 4) if H else 0)


def _put(ws, rr, c, v, fmt=None, fill=WHITE, align=None):
    cell = ws.cell(rr, c, v)
    cell.font = F_BODY
    cell.fill = fill
    if fmt:
        cell.number_format = fmt
    cell.alignment = align or (RIGHT if fmt else LEFT)
    return cell


def _write_values(ws, rr, vals):
    """Columns H..V of a leaf or total row, with the export's red flags."""
    d = _derived(vals)
    seq = [('H', MONEY), ('I', MONEY), ('J', MONEY), ('K', MONEY), ('L', MONEY), ('M', MONEY),
           ('N', MONEY), ('O', MONEY), ('P', PCT1), ('Q', PCT2), ('R', MONEY), ('S', PCT2),
           ('T', MONEY), ('U', MONEY), ('V', MONEY)]
    for i, (k, fmt) in enumerate(seq):
        v = vals[k] if k in vals else d[k]
        v = round(v, 4 if fmt != MONEY else 2)
        red = (k in ('P', 'Q') and v >= 0.7) or (k in ('R', 'S') and v < 0)
        _put(ws, rr, 8 + i, v, fmt, RED if red else WHITE)


def _stamp(wb, snap):
    # fixed document dates so a re-run produces the same file content
    wb.properties.creator = COMPANY
    wb.properties.created = wb.properties.modified = dt.datetime.combine(snap['report'], dt.time(6, 0))


def write_report(path, snap, active_rows, other_rows, currency, entity, groups=True):
    wb = Workbook()
    _stamp(wb, snap)
    ws = wb.active
    ws.title = 'FP&A & QC Project Report'
    ws.sheet_view.showGridLines = False
    rep = snap['report']
    for r, v in ((1, '%s, LLC' % COMPANY), (2, 'FP&A & QC Project Report (In %s)' % currency),
                 (3, 'As of %s' % long_date(rep))):
        ws.cell(r, 1, v).font = F_TITLE
    for r, (a, b) in enumerate((('Reporting Book:', 'ACCRUAL'), ('As of Date:', fmt_date(rep)),
                                ('Location:', entity), ('Project:', 'All Projects')), start=4):
        ws.cell(r, 1, a).font = F_BODY
        ws.cell(r, 2, b).font = F_BODY
    ws.cell(8, 1, ' ').font = F_BODY
    g9 = {8: 'Contract Value', 9: 'Revenue', 10: 'Invoiced Amount', 11: 'Total Revised Budget',
          12: 'Original Budget', 13: 'Total Budget Adjustment', 14: 'Total Costs', 15: ' ',
          16: 'Recognized %', 17: ' ', 18: ' ', 19: ' ', 20: 'Total Costs', 21: 'Total Costs',
          22: 'Invoiced Amount'}
    s10 = {8: 'Inception To Date', 9: 'Inception To Date', 10: 'Inception To Date',
           11: 'Orig Budget + Adjustments', 12: 'Original', 13: 'ETC', 14: 'Inception To Date',
           15: 'Inception To Date', 16: 'REV REC %', 17: 'Budget Spent %', 18: 'Gross profit',
           19: 'Margin %', 20: 'Year To Date', 21: 'Month Ending', 22: 'Year To Date'}
    d11 = {c: fmt_date(rep) for c in range(8, 17)}
    d11.update({17: ' ', 18: ' ', 19: ' ', 20: fmt_date(rep), 21: fmt_date(month_end(rep)), 22: fmt_date(rep)})
    for r, spec in ((9, g9), (10, s10), (11, d11)):
        for c in (1, 2, 3):
            _put(ws, r, c, ' ')
        for c, t in spec.items():
            _put(ws, r, c, t, align=Alignment(horizontal='center'))
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    labels = [' ', 'Dimension name', 'Project status', 'Project Billing type', 'Location Name',
              'Project Manager Name', 'On-Track As of Date', 'EOP', 'Actual', 'Actual', 'Actual',
              'Actual', 'Actual', 'Actual', 'Remaining Budget', 'Actual', 'Summary', 'Summary',
              'Summary', 'Actual', 'Actual', 'Actual']
    for c, t in enumerate(labels, 1):
        _put(ws, 12, c, t)
    rr = 13

    def group(label, indent):
        nonlocal rr
        pad = ' ' * indent
        _put(ws, rr, 1, pad + label)
        _put(ws, rr, 2, pad + label)
        for c in range(3, 23):
            _put(ws, rr, c, pad)
        rr += 1

    def leaves(rows):
        nonlocal rr
        for r in rows:
            for c, v in enumerate([' ' * 6 + r['id'], ' ' * 6 + r['name'], r['status'], r['billing'],
                                   r['location'], r['pm'], r['ot']], 1):
                _put(ws, rr, c, v)
            _write_values(ws, rr, r)
            rr += 1

    def total(label, indent, rows):
        nonlocal rr
        pad = ' ' * indent
        _put(ws, rr, 1, pad + label)
        _put(ws, rr, 2, pad + label)
        for c in range(3, 8):
            _put(ws, rr, c, ' ')
        sums = {k: sum(r[k] for r in rows) for k in 'HIJKLMNTUV'}
        _write_values(ws, rr, sums)
        rr += 1

    group('All Projects', 2)
    if groups:
        group('Active Projects', 4)
        leaves(active_rows)
        total('Total Active Projects', 4, active_rows)
    group('All Other Projects', 4)
    leaves(other_rows)
    total('Total All Other Projects', 4, other_rows)
    total('All Projects', 2, (active_rows if groups else []) + other_rows)
    for _ in range(2):
        _put(ws, rr, 1, ' ')
        rr += 1
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[ws.cell(12, i).column_letter].width = w
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def to_sar(r):
    r = dict(r)
    for k in 'HIJKLMNTUV':
        r[k] = round(r[k] * SAR_PER_USD, 2)
    return r


def write_legacy(path, snap, rows):
    """The pre-08/2026 layout: one flat, pre-filtered 'All Active Projects' tab
    (header rows 8/9/10/11, data from row 12, contract in column G)."""
    wb = Workbook()
    _stamp(wb, snap)
    ws = wb.active
    ws.title = 'All Active Projects'
    ws.sheet_view.showGridLines = False
    ws['A1'] = '%s, LLC' % COMPANY
    ws['A2'] = 'FP&A & QC Project Report'
    ws['A3'] = 'As of %s' % long_date(snap['report'])
    for c in ('A1', 'A2', 'A3'):
        ws[c].font = F_TITLE
    grp = {7: 'Contract Value', 8: 'Revenue', 9: 'Invoiced Amount', 10: 'Total Revised Budget',
           11: 'Total Budget', 12: 'Total Budget Adjustment', 13: 'Total Costs', 17: 'Recognized %'}
    sub = {7: 'Inception To Date', 8: 'Inception To Date', 9: 'Inception To Date',
           10: 'Orig Budget + Adjustments', 11: 'Original', 12: 'ETC', 13: 'Inception To Date',
           14: 'Remaining Budget', 15: 'Orig Budget Remaining', 17: 'REV REC %', 18: 'Budget Spent %'}
    for c, t in grp.items():
        ws.cell(8, c, t).font = F_BODY
    for c, t in sub.items():
        ws.cell(9, c, t).font = F_BODY
    for c in range(7, 18):
        ws.cell(10, c, fmt_date(snap['report'])).font = F_BODY
    labels = ['Project ID', 'Dimension name', 'Project status', 'Project Billing type', 'Location Name',
              'Project Manager Name', 'EOP', 'Actual', 'Actual', 'Actual', 'Actual', 'Actual', 'Actual',
              'Remaining Budget', 'Original Budget Remaining', 'teamwork Status', 'Actual', 'Summary']
    for c, t in enumerate(labels, 1):
        ws.cell(11, c, t).font = F_BODY
    rr = 12
    for r in rows:
        d = _derived(r)
        vals = [r['id'], r['name'], r['status'], r['billing'], r['location'], r['pm'],
                r['H'], r['I'], r['J'], r['K'], r['L'], r['M'], r['N'], d['O'], r['L'] - r['N'],
                None, d['P'], d['Q']]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(rr, c, round(v, 4 if c > 16 else 2) if isinstance(v, float) else v)
            cell.font = F_BODY
            if 7 <= c <= 15:
                cell.number_format = MONEY
            elif c >= 17:
                cell.number_format = PCT2
        rr += 1
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


# ----------------------------------------------------------------------------- main
def self_check(W, L):
    """The current-week roll-up must reproduce the world's own project figures."""
    worst = 0.0
    for p in W.projects:
        f = figures(p, L, CUR, p['status'])
        worst = max(worst, abs(f['cost'] - p['cost_to_date']), abs(f['revenue'] - p['revenue']),
                    abs(f['invoiced'] - p['invoiced']))
    assert worst < 0.05, 'roll-up drifted from the world by $%.2f' % worst


def generate():
    W = world()
    L = Ledger(W)
    self_check(W, L)
    # this week: USD report + E3000 companion
    main, comp, roles = build_rows(W, L, CUR, with_companion=True)
    write_report(CUR_FILE, CUR, pending_rows(W, CUR), main, 'USD', COMPANY.split()[0])
    write_report(E3000_FILE, CUR, [], [to_sar(r) for r in comp], 'SAR', '%s Gulf' % COMPANY.split()[0],
                 groups=False)
    # last week: USD report only (the Gulf entity's report starts with the 09/29 package)
    pmain, _, _ = build_rows(W, L, PREV, with_companion=False)
    write_report(PREV_FILE, PREV, pending_rows(W, PREV), pmain, 'USD', COMPANY.split()[0])
    # archived June workbook in the old layout: pre-filtered to active, commercially-based jobs
    fmain, _, _ = build_rows(W, L, FIX, with_companion=False)
    legacy = [r for r in fmain if r['status'] in ACTIVE and 'OVHD' not in r['id']
              and r['H'] > 0 and (r['L'] > 0 or r['K'] > 0)]
    write_legacy(FIX_FILE, FIX, legacy)
    print('this week : %d rows (+%d pending) -> %s' % (len(main), len(pending_rows(W, CUR)), CUR_FILE))
    print('E3000     : %d rows (dual-listed %s, E3000-only %s, missing %s) -> %s'
          % (len(comp), roles['dual'], roles['only'], sorted(roles['missing']), E3000_FILE))
    print('last week : %d rows -> %s' % (len(pmain), PREV_FILE))
    print('fixture   : %d rows -> %s' % (len(legacy), FIX_FILE))


# Documented regression numbers for demo_data/fixtures/budget_regression.xlsx
# (total rows, projects with an original budget, >70%, 90% & >$50k, 100% Rev Rec FF,
#  original-budget weighted overrun %, revised-ETC weighted overrun %).
EXPECTED = dict(total=717, orig_n=717, gt70=228, gt90=76, ff100=58,
                orig_weighted=1.3558, rev_weighted=0.9813)


def check():
    sys.path.insert(0, os.path.join(ROOT, 'pipeline'))
    import build_overrun_dashboard as b
    c = b.compute(FIX_FILE)
    got = dict(total=c['total'], orig_n=c['orig_n'], gt70=c['gt70'], gt90=c['gt90'], ff100=c['ff100'],
               orig_weighted=round(c['orig_weighted'], 4), rev_weighted=round(c['rev_weighted'], 4))
    print('fixture computes:', got)
    assert got == EXPECTED, 'regression FAILED: expected %s' % EXPECTED
    print('regression PASSED')


if __name__ == '__main__':
    if '--check' in sys.argv:
        check()
    else:
        generate()
