"""
sync_util_data.py
=================
Parses index.html and writes util_data.js
to the same folder. util_data.js exposes window.CH_UTIL_DATA used by
location-scorecards.html for live utilization targets/actuals,
non-posting hours, and monthly billable hours by office.

Also reads monthly billable hours (excluding Overhead Non-Posting tasks)
from the source Excel workbook when it can be located alongside this script.

Run manually or via scheduled Cowork task whenever the Utilization Site
or Excel file is updated.
"""

import re
import json
import os
import glob
from datetime import datetime, timezone

# ── Paths (relative to this script) ─────────────────────────────────────────
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SOURCE_HTML = os.path.join(SCRIPT_DIR, 'index.html')
OUTPUT_JS   = os.path.join(SCRIPT_DIR, 'util_data.js')

# ── Month col mapping in PMO_Monthly table ───────────────────────────────────
# Each month has 3 columns (Tgt, Act, Δ). The data-col attr identifies the group.
MONTH_COLS = [1, 4, 7, 10, 13, 16, 19, 22, 25, 28, 31, 34]   # Jan=1 … Dec=34

# ── Helpers ──────────────────────────────────────────────────────────────────

def parse_pct(s):
    """'81%' → 0.81, 'N/A' or '—' or '·' → None"""
    s = s.strip()
    if not s or s in ('N/A', '—', '·', '-'):
        return None
    s = s.replace('%', '').replace(',', '').strip()
    try:
        return round(float(s) / 100, 4)
    except ValueError:
        return None

def parse_num(s):
    """'232' → 232, '1,826' → 1826, '·' → 0"""
    s = s.strip()
    if not s or s in ('—', '-', 'N/A'):
        return None
    if s == '·':
        return 0
    s = s.replace(',', '').strip()
    try:
        return int(float(s))
    except ValueError:
        return None

def strip_tags(s):
    return re.sub(r'<[^>]+>', '', s).strip()

# ── 1. Parse PMO_Monthly ─────────────────────────────────────────────────────

def parse_pmo_monthly(html):
    """
    Returns:
        util_targets: {office: [12 floats or None]}   Jan-Dec targets
        util_actuals: {office: [12 floats or None]}   Jan-Dec actuals (future = None)
    """
    # Locate section
    idx = html.find('id="PMO_Monthly"')
    if idx == -1:
        raise RuntimeError('PMO_Monthly section not found in HTML')
    section = html[idx:idx + 200000]

    # Find first tbody
    tb_start = section.find('<tbody>')
    tb_end   = section.find('</tbody>', tb_start) + len('</tbody>')
    if tb_start == -1:
        raise RuntimeError('No tbody found in PMO_Monthly section')
    tbody = section[tb_start:tb_end]

    # Split into rows
    rows = re.findall(r'<tr[^>]*>(.*?)</tr>', tbody, re.DOTALL)

    util_targets = {}
    util_actuals = {}

    for row in rows:
        # Office name
        m = re.search(r'<td class="col-office"[^>]*>(.*?)</td>', row, re.DOTALL)
        if not m:
            continue
        office = strip_tags(m.group(1)).strip()
        if not office or office.lower() in ('total', 'pmo'):
            continue
        # Skip footnote rows (start with * or are clearly not office names)
        if office.startswith('*') or len(office) > 80:
            continue

        targets = []
        actuals = []

        for col in MONTH_COLS:
            # Target: td with class containing "period-sep" and data-col=col
            tgt_pat = (
                r'<td class="num period-sep[^"]*"\s+data-col="' + str(col) + r'">'
                r'([^<]+)</td>'
            )
            tm = re.search(tgt_pat, row)
            targets.append(parse_pct(tm.group(1)) if tm else None)

            # Actual: td with data-col=col that contains a <span class="pill ...">
            act_pat = (
                r'<td class="num[^"]*"\s+data-col="' + str(col) + r'">'
                r'<span class="pill[^"]*">([^<]+)</span>'
            )
            am = re.search(act_pat, row)
            if am:
                val = am.group(1)
                actuals.append(None if val == 'N/A' else parse_pct(val))
            else:
                actuals.append(None)

        util_targets[office] = targets
        util_actuals[office] = actuals

    return util_targets, util_actuals

# ── 2. Parse Non-Posting Hours ───────────────────────────────────────────────

def parse_nonpost_hours(html):
    """
    Returns:
        nonpost_hrs: {office: [Jan, Feb, Mar, Apr, May, Total]}  raw integer hours
    """
    idx = html.find('id="Summary_Non_Posting"')
    if idx == -1:
        raise RuntimeError('Summary_Non_Posting section not found in HTML')
    section = html[idx:idx + 100000]

    # There are two tbodies: first = percentages, second = hours
    tbodies = re.findall(r'<tbody>(.*?)</tbody>', section, re.DOTALL)
    if len(tbodies) < 2:
        raise RuntimeError('Expected 2 tbodies in Summary_Non_Posting, found ' + str(len(tbodies)))

    hours_tbody = tbodies[1]
    rows = re.findall(r'<tr[^>]*>(.*?)</tr>', hours_tbody, re.DOTALL)

    nonpost_hrs = {}

    for row in rows:
        # Office name in sticky-l td
        m = re.search(r'<td class="sticky-l[^"]*"[^>]*>(.*?)</td>', row, re.DOTALL)
        if not m:
            continue
        office = strip_tags(m.group(1)).strip()
        if not office or office.lower() == 'total':
            continue

        # All numeric td cells (exclude sticky-l)
        cells = re.findall(
            r'<td class="num[^"]*"[^>]*>(.*?)</td>',
            row, re.DOTALL
        )
        vals = [parse_num(strip_tags(c)) for c in cells]

        # Expect 6 values: Jan, Feb, Mar, Apr, May, Total
        # Pad with None if shorter
        while len(vals) < 6:
            vals.append(None)
        nonpost_hrs[office] = vals[:6]

    return nonpost_hrs

# ── 3. Billable Hours from Excel ─────────────────────────────────────────────

# Maps Raw Data location names → scorecard names where they differ
RAW_TO_SCORECARD = {
    'NH Gulf (USD)':          'NH Gulf',     # the International book's ERP location
    'International':          'NH Gulf',
    'Global LLC International':'NH Gulf Global LLC',
}

# The fiscal year being reported: FPA_AS_OF=YYYY-MM-DD pins it for a past package, otherwise
# the current calendar year (unchanged behaviour).
REPORT_YEAR = int((os.environ.get('FPA_AS_OF') or datetime.now().strftime('%Y-%m-%d'))[:4])

# Task names that represent Overhead Non-Posting (excluded from billable hours)
OVHD_NONPOST_TASKS = {'Overhead Non-Posting', 'CA - Overhead Non-Posting'}

def find_excel(script_dir):
    """Return path to the Utilization Report xlsx, or None if not found."""
    # The workbook now normally lives in the Weekly and Monthly Updates\Utilization
    # subfolder; search there before falling back to a bare *.xlsx in this folder
    # (that catch-all used to pick up unrelated workbooks such as
    # "Projects 70+% of Budget.xlsx", which has no Raw Data sheet, silently
    # emptying billable_hrs).
    patterns = [
        os.path.join(script_dir, 'Utilization Report*.xlsx'),
        os.path.join(script_dir, 'utilization_report*.xlsx'),
        os.path.join(script_dir, 'Weekly and Monthly Updates', 'Utilization', 'Utilization Report*.xlsx'),
        os.path.join(script_dir, 'Weekly and Monthly Updates', '**', 'Utilization Report*.xlsx'),
        os.path.join(script_dir, '*.xlsx'),
    ]
    for pat in patterns:
        matches = [p for p in glob.glob(pat, recursive=True)
                   if not os.path.basename(p).startswith('~$')]
        if matches:
            return matches[0]
    return None

def parse_billable_hours(xlsx_path):
    """
    Read Raw Data sheet from the Excel workbook and return:
        billable_hrs: {office: [Jan, Feb, Mar, Apr, May, Jun, Jul, Aug, Sep, Oct, Nov, Dec]}
    Hours are summed from rows where Task name is NOT in OVHD_NONPOST_TASKS
    and Utilized hours > 0.
    """
    try:
        import openpyxl
    except ImportError:
        print('  openpyxl not installed — skipping billable hours extraction')
        return {}

    print(f'  Reading: {xlsx_path}')
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    if 'Raw Data' not in wb.sheetnames:
        print('  Raw Data sheet not found in workbook — skipping')
        return {}

    ws = wb['Raw Data']
    from collections import defaultdict
    acc = defaultdict(lambda: [0.0] * 12)

    for row in ws.iter_rows(min_row=2, values_only=True):
        task     = str(row[10]).strip() if row[10] else ''
        if task in OVHD_NONPOST_TASKS:
            continue
        loc      = row[2]
        month_dt = row[25]   # datetime with last-day-of-month
        util_hrs = row[5]    # Utilized hours
        if not loc or not isinstance(month_dt, datetime) or not isinstance(util_hrs, (int, float)):
            continue
        if util_hrs <= 0:
            continue
        if month_dt.year != REPORT_YEAR:   # exclude prior-year carryover (e.g. last December)
            continue
        loc = RAW_TO_SCORECARD.get(loc, loc)
        acc[loc][month_dt.month - 1] += util_hrs

    # Round to whole numbers
    return {loc: [round(h) for h in hrs] for loc, hrs in acc.items()}


# ── 4. Main ──────────────────────────────────────────────────────────────────

def main():
    print(f'Reading: {SOURCE_HTML}')
    with open(SOURCE_HTML, 'r', encoding='utf-8') as f:
        html = f.read()

    print('Parsing PMO_Monthly…')
    util_targets, util_actuals = parse_pmo_monthly(html)
    print(f'  → {len(util_targets)} offices found')

    print('Parsing Non-Posting Hours…')
    nonpost_hrs = parse_nonpost_hours(html)
    print(f'  → {len(nonpost_hrs)} offices found')

    # Billable hours from Excel (optional)
    print('Looking for Excel workbook…')
    xlsx_path = find_excel(SCRIPT_DIR)
    billable_hrs = {}
    if xlsx_path:
        print(f'  Found: {os.path.basename(xlsx_path)}')
        print('Parsing Billable Hours (Raw Data, excl. Overhead Non-Posting)…')
        billable_hrs = parse_billable_hours(xlsx_path)
        print(f'  → {len(billable_hrs)} offices found')
    else:
        print('  No Excel file found — billable_hrs will be empty in output')

    # Build output object
    data = {
        'generated':    datetime.now(timezone.utc).isoformat(),
        'source_file':  'index.html',
        'util_targets': util_targets,
        'util_actuals': util_actuals,
        'nonpost_hrs':  nonpost_hrs,
        'billable_hrs': billable_hrs,
    }

    js = (
        '// Auto-generated by sync_util_data.py — do not edit manually.\n'
        '// Re-run sync_util_data.py to refresh from index.html\n'
        'window.CH_UTIL_DATA = '
        + json.dumps(data, indent=2, ensure_ascii=False)
        + ';\n'
    )

    with open(OUTPUT_JS, 'w', encoding='utf-8') as f:
        f.write(js)

    print(f'Written: {OUTPUT_JS}')

    # Spot-check one office (the first on the PMO Monthly table)
    spot = next(iter(util_targets), '')
    print(f'\n  {spot} targets (12 mo):      {util_targets.get(spot, [])}')
    print(f'  {spot} actuals (12 mo):      {util_actuals.get(spot, [])}')
    print(f'  {spot} non-post hrs   :      {nonpost_hrs.get(spot, [])}')
    print(f'  {spot} billable hrs (12 mo): {billable_hrs.get(spot, [])}')
    print(f'  billable-hours offices: {len(billable_hrs)}')

if __name__ == '__main__':
    main()
