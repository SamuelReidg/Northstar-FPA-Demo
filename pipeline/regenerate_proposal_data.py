"""
Regenerate proposal_data.js from the source "Proposal Hours Summary" export.

Usage:
    python regenerate_proposal_data.py [<source.xlsx>]

If no path is provided, defaults to the most recent
"*Proposal Hours Summary*.xlsx" in the same folder.

This powers the "proposal-command-center.html" dashboard. It reads the
summary tabs of the workbook (pivot-style summaries FP&A maintains) and
writes a compact proposal_data.js next to itself.

Monthly refresh workflow (mirrors the PM Hours dashboard):
  1. FP&A drops the new "Proposal Hours Summary_<date>.xlsx" into the package folder.
  2. Run this script -> proposal_data.js.
  3. Copy proposal_data.js next to proposal-command-center.html on the site share.
  4. Hard-refresh the dashboard (Ctrl+F5).

Required packages: openpyxl
"""

from __future__ import annotations
import json
import sys
import datetime as dt
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("ERROR: openpyxl is required. Run:  pip install openpyxl")

# ----------------------------------------------------------------------
# Sheet names (update only if the workbook template is renamed)
# ----------------------------------------------------------------------
SHEET_OFFICE    = "Office Summary (Proj. Loc.)"
SHEET_SEGMENT   = "Segmented Summary (Proj. Loc.)"
SHEET_AMOUNT    = "$ Amount"
SHEET_RECORDS   = "# Records"

MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------
def pn(v):
    """Parse a numeric cell -> float, or None for #N/A, '-', '', #DIV/0! etc."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip()
    if s in ("", "-", "N/T", "#N/A", "#DIV/0!", "#REF!", "#VALUE!"):
        return None
    try:
        return float(s.replace(",", "").replace("$", ""))
    except ValueError:
        return None


def month_label(d):
    """datetime -> "Jul '25" style label."""
    if isinstance(d, dt.datetime):
        return f"{MONTH_ABBR[d.month - 1]} '{d.year % 100:02d}"
    return str(d)


def find_month_cols(ws, header_row, first_col=2, last_col=14):
    """Return (col_indices, labels) for date-typed header cells in a row."""
    cols, labels = [], []
    for c in range(first_col, last_col + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, dt.datetime):
            cols.append(c)
            labels.append(month_label(v))
    return cols, labels


def date_blocks(ws, header_row, first_col=2, last_col=None):
    """Return a list of contiguous runs of date-typed header cells.

    Each run is (cols, labels).  Layout-agnostic: survives the workbook
    gaining or losing month columns (added 2026-08-18 after the July export
    grew from 12 to 13 months and silently truncated every series)."""
    if last_col is None:
        last_col = ws.max_column
    runs, cur, curl = [], [], []
    for c in range(first_col, last_col + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, dt.datetime):
            cur.append(c); curl.append(month_label(v))
        elif cur:
            runs.append((cur, curl)); cur, curl = [], []
    if cur:
        runs.append((cur, curl))
    return runs


def col_after_block(ws, header_row, cols):
    """Column index of the first labelled (non-blank) header right of a block."""
    if not cols:
        return None
    for c in range(max(cols) + 1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is not None and str(v).strip() != "":
            return c
    return None


def row_by_label(ws, label, col=1, rmax=60):
    """Find the 1-based row index whose `col` cell trims to `label`."""
    target = label.strip().lower()
    for r in range(1, rmax + 1):
        v = ws.cell(row=r, column=col).value
        if v is not None and str(v).strip().lower() == target:
            return r
    return None


# ----------------------------------------------------------------------
# Main extraction
# ----------------------------------------------------------------------
def build(src_path: Path) -> dict:
    wb = load_workbook(src_path, read_only=True, data_only=True)

    # ---- Office Summary: months + per-office hours & proposal size --------
    ws = wb[SHEET_OFFICE]
    _blocks = date_blocks(ws, 2)
    hours_cols, month_labels = _blocks[0] if len(_blocks) > 0 else ([], [])
    size_cols, _ = _blocks[1] if len(_blocks) > 1 else ([], [])
    HOURS_AVG_COL = col_after_block(ws, 2, hours_cols)
    SIZE_AVG_COL = col_after_block(ws, 2, size_cols)

    offices = []
    grand_size_trend = []
    for r in range(3, 60):
        name = ws.cell(row=r, column=1).value
        if name is None:
            continue
        name = str(name).strip()
        low = name.lower()
        if low.startswith("grand total"):
            grand_size_trend = [pn(ws.cell(row=r, column=c).value) for c in size_cols]
            break  # end of the first per-opportunity office block
        if (low.endswith("total") or low.endswith("total**") or low.endswith("*")
                or low.startswith("corporate") or low.startswith("sales and")
                or low.startswith("<") or low.startswith("8 -") or low.startswith("> ")):
            continue
        hours = [pn(ws.cell(row=r, column=c).value) for c in hours_cols]
        if all(h is None for h in hours):
            continue
        avg = pn(ws.cell(row=r, column=HOURS_AVG_COL).value)
        size = [pn(ws.cell(row=r, column=c).value) for c in size_cols]
        size_avg = pn(ws.cell(row=r, column=SIZE_AVG_COL).value)
        offices.append({
            "name": name,
            "hours": hours,
            "avg": avg,
            "size": size,
            "size_avg": size_avg,
        })

    # ---- Segmented Summary: region hours-per-opp, hours, proposals --------
    ws = wb[SHEET_SEGMENT]
    _segb = date_blocks(ws, 2)
    seg_cols, _ = _segb[0] if _segb else ([], [])

    def seg_series(label):
        r = row_by_label(ws, label, col=1, rmax=40)
        if r is None:
            return None
        return [pn(ws.cell(row=r, column=c).value) for c in seg_cols]

    region_hpo = {
        "Domestic Roll-Up": seg_series("Domestic Roll-Up"),
        "Meridian": seg_series("Meridian"),
        "International": seg_series("International"),
        "Grand Total": seg_series("Grand Total"),
    }

    def block_totals(block_title):
        hr = row_by_label(ws, block_title, col=1, rmax=40)
        out = {}
        if hr is None:
            return out
        total_col = None
        for c in range(2, ws.max_column + 1):
            if str(ws.cell(row=hr + 1, column=c).value).strip().lower() == "total":
                total_col = c
                break
        for r in range(hr + 2, hr + 8):
            lbl = ws.cell(row=r, column=1).value
            if lbl is None:
                continue
            lbl = str(lbl).strip()
            out[lbl] = pn(ws.cell(row=r, column=total_col).value) if total_col else None
            if lbl.lower().startswith("grand total"):
                break
        return out

    proposals_by_region = block_totals("# of Proposals")
    hours_by_region = block_totals("Hours Spent (Proposal Time)")

    # ---- Monthly region series: $ value and average proposal size ---------
    REGIONS = ["Domestic", "Meridian", "International", "Grand Total"]

    def region_block_series(block_title):
        hr = row_by_label(ws, block_title, col=1, rmax=ws.max_row)
        out = {}
        if hr is None:
            return out
        for r in range(hr + 1, hr + 9):
            lbl = ws.cell(row=r, column=1).value
            if lbl is None:
                continue
            lbl = str(lbl).strip()
            if lbl in REGIONS:
                out[lbl] = [pn(ws.cell(row=r, column=c).value) for c in seg_cols]
            if lbl.lower().startswith("grand total"):
                break
        return out

    value_by_region = region_block_series("$ of Proposals")
    avg_size_by_region = region_block_series("Average Proposal Size by Region")

    # ---- KPI totals from $ Amount and # Records "Total" rows --------------
    def total_row_sum(sheet_name):
        w = wb[sheet_name]
        for r in range(1, w.max_row + 1):
            if str(w.cell(row=r, column=2).value).strip().lower() == "total":
                vals = [pn(w.cell(row=r, column=c).value)
                        for c in range(3, w.max_column + 1)]
                vals = [v for v in vals if v is not None]
                return sum(vals), vals
        return None, []

    total_value, value_monthly = total_row_sum(SHEET_AMOUNT)
    total_opps, opps_monthly = total_row_sum(SHEET_RECORDS)

    # ---- Hours distribution: bucket every office-month cell ---------------
    lt8 = b810 = gt10 = 0
    for o in offices:
        for h in o["hours"]:
            if h is None:
                continue
            if h < 8:
                lt8 += 1
            elif h <= 10:
                b810 += 1
            else:
                gt10 += 1
    dist_total = max(lt8 + b810 + gt10, 1)

    # ---- KPIs -------------------------------------------------------------
    woff = wb[SHEET_OFFICE]
    grg = None
    for _r in range(1, 80):
        _v = woff.cell(row=_r, column=1).value
        if _v is not None and str(_v).strip().lower().startswith("grand total"):
            grg = _r
            break
    grand_avg_hours = pn(woff.cell(row=grg, column=HOURS_AVG_COL).value) if grg else None
    grand_size_avg = pn(woff.cell(row=grg, column=SIZE_AVG_COL).value) if grg else None
    if grand_avg_hours is None:
        grand_avg_hours = (region_hpo["Grand Total"] or [None])[-1]
    if grand_size_avg is None and total_opps:
        grand_size_avg = (total_value / total_opps) if total_value else None

    kpis = {
        "avg_hours_per_opp": grand_avg_hours,
        "avg_proposal_size": grand_size_avg,
        "total_opportunities": total_opps,
        "total_proposal_value": total_value,
        "total_hours": hours_by_region.get("Grand Total"),
        "num_offices": len(offices),
    }

    # ---- Individual Summary (Emp. Loc.): employees grouped by office -------
    def build_individual():
        wsi = wb["Individual Summary (Emp. Loc.)"]
        # Month columns + the pivot's "Grand Total" column are found from the
        # header row (layout-agnostic, like the Office Summary blocks), so the
        # table always shows every month and the real total.
        _ib = date_blocks(wsi, 4)
        mcols2, mlabels2 = _ib[0] if _ib else ([], [])
        TOTAL_COL = col_after_block(wsi, 4, mcols2) or 13
        offs, cur, grand_ind = [], None, None
        for r in range(5, wsi.max_row + 1):
            a = wsi.cell(row=r, column=1).value
            if a is None:
                continue
            a = str(a).strip()
            months = [pn(wsi.cell(row=r, column=c).value) for c in mcols2]
            total = pn(wsi.cell(row=r, column=TOTAL_COL).value)
            if a.lower().startswith("grand total"):
                grand_ind = {"months": months, "total": total}
                break
            if "," in a:                       # employee row
                if cur is not None:
                    cur["employees"].append({"name": a, "months": months, "total": total})
            else:                              # office subtotal row
                if a == "0":
                    a = "Unassigned"
                cur = {"office": a, "months": months, "total": total, "employees": []}
                offs.append(cur)
        return {"months": mlabels2, "offices": offs, "grand": grand_ind}

    individual = build_individual()

    return {
        "meta": {
            "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "source_file": src_path.name,
            "months": month_labels,
            "latest_month": month_labels[-1] if month_labels else None,
        },
        "kpis": kpis,
        "offices": offices,
        "individual": individual,
        "region_hours_per_opp": region_hpo,
        "proposals_by_region": proposals_by_region,
        "hours_by_region": hours_by_region,
        "value_monthly": value_monthly,
        "opps_monthly": opps_monthly,
        "avg_size_trend": grand_size_trend,
        "value_by_region": value_by_region,
        "avg_size_by_region": avg_size_by_region,
        "hours_distribution": {
            "lt8": lt8, "b810": b810, "gt10": gt10,
            "lt8_pct": round(100 * lt8 / dist_total),
            "b810_pct": round(100 * b810 / dist_total),
            "gt10_pct": round(100 * gt10 / dist_total),
        },
    }


def main():
    here = Path(__file__).resolve().parent
    if len(sys.argv) > 1:
        src = Path(sys.argv[1])
    else:
        cands = sorted(here.glob("*Proposal Hours Summary*.xlsx"),
                       key=lambda p: p.stat().st_mtime, reverse=True)
        if not cands:
            sys.exit("ERROR: no '*Proposal Hours Summary*.xlsx' found. Pass a path.")
        src = cands[0]

    print(f"Reading: {src}")
    data = build(src)
    out = here / "proposal_data.js"
    js = "// Auto-generated by regenerate_proposal_data.py - do not edit by hand.\n"
    js += "window.PROPOSAL_DATA = " + json.dumps(data, indent=1) + ";\n"
    out.write_text(js, encoding="utf-8")
    k = data["kpis"]
    print(f"Wrote {out}  ({out.stat().st_size/1024:.1f} KB)")
    print(f"  months: {data['meta']['months']}")
    print(f"  offices: {k['num_offices']}")
    print(f"  avg hrs/opp: {k['avg_hours_per_opp']}")
    print(f"  avg proposal size: {k['avg_proposal_size']}")
    print(f"  total opportunities: {k['total_opportunities']}")
    print(f"  total proposal value: {k['total_proposal_value']}")
    print(f"  total hours: {k['total_hours']}")
    print(f"  distribution: {data['hours_distribution']}")
    print(f"  proposals by region: {data['proposals_by_region']}")


if __name__ == "__main__":
    main()
