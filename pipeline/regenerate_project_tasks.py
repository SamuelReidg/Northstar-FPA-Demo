#!/usr/bin/env python3
"""regenerate_project_tasks.py -- builds project_tasks.js (window.PROJECT_TASKS)
for the "Task Budget Detail" panel in project-recovery.html.

Scope = the SAME at-risk, In Progress set the EWS dashboard shows
(Critical + Warning + Watch), so every project on the Projects tab has task
detail regardless of tier. Source = the monthly "SR of Projects @ Risk
<date>.xlsx":
  - sheet 'Project Details ' (trailing space): determines the at-risk set
    using the EWS tier rule (worse-of cost%/hours%, x100; >=5 = at-risk).
  - sheet 'Budget to Actual Costs by Task': per-task budget vs actual rows.

Each task row written as [name, budgeted$, spent$, overrun$, budgetHrs, actualHrs, overHrs]
matching the schema project-recovery.html expects.

Usage: python regenerate_project_tasks.py "SR of Projects @ Risk <date>.xlsx" [out.js]
"""
import sys, json, re
from collections import defaultdict
import openpyxl

PD_SHEET = "Project Details "          # note trailing space
TASK_SHEET = "Budget to Actual Costs by Task"
PID_RE = re.compile(r"^\d{2}-[A-Z]{2}-\d{5}$")

def num(x):
    try: return float(x)
    except (TypeError, ValueError): return 0.0

def clean_name(s):
    s = str(s or "").strip()
    # source names are "<task code> - <description>"; drop the leading code segment
    return s.split(" - ", 1)[1].strip() if " - " in s else s

def build_atrisk(wb):
    ws = wb[PD_SHEET]
    atrisk = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        pid = str(r[0] or "").strip()
        if not pid: continue
        status = str(r[3] or "").strip().lower()
        cost = num(r[10]) * 100      # Task Overrun %
        hrs  = num(r[16]) * 100      # Hours Overrun %
        m = max(cost, hrs)
        if status == "in progress" and m >= 5:
            atrisk[pid] = "C" if m >= 25 else "W" if m >= 10 else "T"
    return atrisk

def collect_tasks(wb, atrisk):
    """Read the task tab layout-agnostically.

    The 08/2026 ERP export moved the header row (11 -> 12), inserted a spacer
    column that shifted every money column right one, and DROPPED the
    "Task Overrun $" / "Task Overrun Hours" columns.  The automated workbook's
    hidden `_Calc` tab recomputes both, row-for-row against this tab, so we
    read the overrun figures from there and stay consistent with
    'Project Details '.  Requires a workbook whose formulas have been
    calculated (open+save in Excel, or a LibreOffice headless recalc).
    """
    ws = wb[TASK_SHEET]
    grid = list(ws.iter_rows(values_only=True))

    hdr = None
    for i, r in enumerate(grid):
        if any(str(c or "").strip() == "Task Project ID" for c in r[:6]):
            hdr = i
            break
    if hdr is None:
        raise SystemExit("ERROR: no 'Task Project ID' header found on '%s'" % TASK_SHEET)

    def col(*names, required=True):
        for j, c in enumerate(grid[hdr]):
            if str(c or "").strip() in names:
                return j
        if required:
            raise SystemExit("ERROR: column %s not found on '%s'" % (names, TASK_SHEET))
        return None

    C_PID  = col("Task Project ID")
    C_REV  = col("Revised Budget", "Revised Budget ")
    C_BHRS = col("Budgeted Hours")
    C_AHRS = col("Actual Hours")
    # "Total Costs" arrives as a bare "Actual" sub-header; it is the last
    # money column before Budgeted Hours.
    C_COST = C_BHRS - 1

    calc = {}
    if "_Calc" in wb.sheetnames:
        for i, r in enumerate(wb["_Calc"].iter_rows(values_only=True)):
            calc[i] = (num(r[3]) if len(r) > 3 else 0.0,
                       num(r[4]) if len(r) > 4 else 0.0)

    def derive(name, revbud, cost, bhrs, ahrs):
        """Fallback overrun rule (mirrors _Calc) when the tab is absent."""
        n = str(name or "").strip()
        if n.endswith("OVHD-Non-Posting - Overhead Non-Posting"):
            return cost, ahrs
        if revbud > 0:
            return max(0.0, cost - revbud), max(0.0, ahrs - bhrs)
        return 0.0, 0.0

    out = defaultdict(list)
    for i in range(hdr + 2, len(grid)):
        r = grid[i]
        pid = str(r[C_PID] or "").strip()
        if pid not in atrisk:
            continue
        revbud, cost = num(r[C_REV]), num(r[C_COST])
        bhrs, ahrs = num(r[C_BHRS]), num(r[C_AHRS])
        if i in calc:
            ovr, ovhrs = calc[i]
        else:
            ovr, ovhrs = derive(r[0], revbud, cost, bhrs, ahrs)
        out[pid].append([
            clean_name(r[0]),
            round(revbud), round(cost), round(ovr),
            round(bhrs, 1), round(ahrs, 1), round(ovhrs, 1),
        ])
    for pid in out:
        out[pid].sort(key=lambda t: -t[3])   # worst overrun $ first
    return out


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "project_tasks.js"
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    atrisk = build_atrisk(wb)
    tasks = collect_tasks(wb, atrisk)
    data = {pid: tasks[pid] for pid in tasks}
    payload = "window.PROJECT_TASKS=" + json.dumps(data, separators=(",", ":")) + ";\n"
    with open(out, "w", encoding="utf-8") as f:
        f.write(payload)
    tiers = {"C":0,"W":0,"T":0}
    for pid in data: tiers[atrisk[pid]] += 1
    print("at-risk projects:", len(atrisk), "| with tasks:", len(data),
          "| tiers C/W/T:", tiers["C"], tiers["W"], tiers["T"])
    print("total task rows:", sum(len(v) for v in data.values()))
    print("bytes:", len(payload.encode("utf-8")))
    print("saved ->", out)

if __name__ == "__main__":
    main()
