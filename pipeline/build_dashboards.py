"""
Rebuild engine for the Approvals Dashboard (Timesheets & Expense Reports,
Draft/Submitted).

SOURCES CHANGED 2026-08-20. Going forward the dashboard is built from THREE files,
not two workbooks:

  1. Timesheets  <n>.csv      one row per timesheet
        Employee ID | Employee name | Employee last name | Employee first name |
        Begin date | End date | GL post date | State        (+ junk Print/Duplicate cols)
  2. Employees <n>.csv        the employee master (replaces the old ERP/HR-system/contractor tabs)
        Employee ID | Legacy ID | Primary contact name | Title | First name | Last name |
        Primary email address | Location ID | Department ID | Manager | Supervisor name |
        FLSA STATUS | PAYGROUP | Salesforce ID | Not Employed
  3. Staff expenses <n>.xlsx  one row per expense report, sheet "Staff expenses"
        Employee ID | Employee Name | ER # | Reason for expense | State | Date Filed |
        GL Posting Date | Amount | Date Paid | Amount Due | Last Modified | Summary State

The three files may be passed in ANY ORDER -- resolve_sources() sniffs them by
content, because the ERP stamps a drifting number into each filename
("Timesheets 9.csv", "Employees 32 1.csv", "Staff expenses 22.xlsx").

WHAT THE OLD SOURCES GAVE US AND THESE DO NOT:
  * TARGET. The old timesheet workbook carried the HR system's 'Target (Text)' and a
    contractor-platform target lookup, which drove a `Target != 0` filter. Neither new file has any
    target column, so THAT FILTER IS GONE. Non-billable/corporate people are now kept
    out by name via EXCLUDE_NAMES in build_dashboard_html.py instead.
  * LOCATION NAMES. The old expense workbook had a Location Data tab mapping
    Location ID -> office name. The new expense file is a single sheet, so the map is
    baked in below as LOCATION_NAMES (harvested from the 2026-08-18 export). An
    unknown ID falls back to showing the raw ID, so a new office shows up as "L8190"
    rather than disappearing from the PMO filter.

Columns are still resolved BY HEADER NAME, never by position, in both the CSVs and
the workbook -- so column shifts are harmless and a genuine rename fails loudly.
"""
import openpyxl, datetime, csv, os, re
from collections import defaultdict

# ---------------------------------------------------------------------------
# WHICH ITEMS THE DASHBOARD SHOWS  (settled 2026-08-20)
#
# An open item (State = Draft or Submitted) is shown only if it has been
# outstanding MORE THAN MIN_DAYS_OUTSTANDING days as of the export's own date.
# BOTH sides age on GL post date / GL Posting Date -- one field, one meaning,
# verified row for row against an independent reference build. Demo package
# 9.29.2026 (as of 2026-09-26): timesheets 30/9/39 over 15 supervisors; expenses
# 43/79/122 over 24 supervisors, $77,986.54; 129 employees affected.
#
# Rule history -- any figure quoted before 2026-08-20 came from a superseded rule:
#   08-18  timesheets kept the latest GL date holding >=10% of open rows, which IS
#          the in-flight week. Listed dozens of sheets that were not late, hid the real
#          ones (avg_age rendered 0.0 -- the tell). Expenses had no date filter.
#   08-19  "past due" = activity before the start of the current week; timesheets
#          aged on End date, expenses on Last Modified.
#   08-20  this ">10 days outstanding on GL" gate, on the new three-file sources.
#
# KNOWN CHARACTERISTIC of GL, read before "fixing" a count: GL post date is when the
# item POSTED, not the week a timesheet covers. Stragglers get batched onto a recent
# GL date, so a sheet covering the week ending 04/11/2026 that posted to GL 09/01 reads
# as 25 days outstanding rather than 168. End date is still carried on every row and shown in the
# detail table -- set AGE_KEYS_TS = ('end', 'gl') if you ever want true age instead.
MIN_DAYS_OUTSTANDING  = 10           # strictly greater than this many days
AGE_KEYS_TS           = ('gl',)      # first present key wins, per row
AGE_KEYS_EX           = ('gl',)      # first present key wins, per row
DAYS_OUTSTANDING_ONLY = True         # False = every open item, no age gate at all
AS_OF                 = None         # None = derive from the export; or datetime.date(Y,M,D)

OPEN_STATES = ('Draft', 'Submitted')  # everything else (Approved/Paid/Declined/...) is closed

# ERP Location ID -> office (PMO) name. Harvested from the Location Data tab of the
# last expense workbook that shipped one. Unknown IDs fall back to the ID -- in the demo
# package L8190 (Geospatial) and L8210 (Meridian Southwest) are deliberately left out so
# the raw-ID fallback is visible in the PMO filter; add them here once the names are known.
LOCATION_NAMES = {
    'L8000': 'Big Sky', 'L8010': 'Bluegrass', 'L8020': 'Cascade', 'L8030': 'Chesapeake',
    'L8040': 'Front Range', 'L8050': 'Great Lakes', 'L8060': 'Gulf Coast',
    'L8070': 'High Desert', 'L8080': 'Island', 'L8090': 'Lakeshore', 'L8100': 'Hill Country',
    'L8110': 'Piedmont', 'L8120': 'Prairie', 'L8130': 'Red Rock', 'L8140': 'Sierra',
    'L8150': 'Coastal Bend', 'L8160': 'Built Environment', 'L8170': 'Paleontology',
    'L8180': 'Maritime', 'L8200': 'Meridian Central', 'L8220': 'NH Gulf (USD)',
    'L8230': 'Corporate', 'L8240': 'Sales and Marketing', 'L8250': 'Archive Services',
    'L7010': 'Riverton - INACTIVE', 'L7020': 'Coastal Plain - INACTIVE',
    'L7030': 'Northern Plains - INACTIVE', 'L7040': 'Canyon Lands - INACTIVE',
}

# The expense workbook's single sheet, with the names it has shipped under.
SHEET_ALIASES = {
    'Staff expenses': ('Staff expenses', 'Staff expenses (1)', 'Expenses Data',
                       'Expenses', 'Sheet1'),
}


# ------------------------------------------------------------------ primitives
def as_num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


_EMP_SUFFIX = re.compile(r'\s*-\s*employee\s*$', re.I)

def clean_name(v):
    """
    Normalise a person's name for display. The ERP appends " - Employee" to some
    names ("Doe, Alex - Employee"), which makes the same person look like two
    people between files and breaks the exclusion list. Strip it.
    """
    if v is None:
        return None
    s = _EMP_SUFFIX.sub('', str(v).strip())
    return s or None


_DATE_FORMATS = ('%m/%d/%Y %I:%M:%S %p', '%m/%d/%Y %H:%M:%S', '%m/%d/%Y',
                 '%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%m/%d/%y')

def as_date(v):
    """Parse a cell/CSV value into a datetime, or None. Times are discarded later."""
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime(v.year, v.month, v.day)
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    for f in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def _header_map(values):
    """Map lower-cased header text -> 0-based column index."""
    m = {}
    for i, v in enumerate(values or ()):
        if v is not None and str(v).strip() != '':
            m.setdefault(str(v).strip().lower(), i)
    return m


def _satisfies(hm, required):
    """required = tuple of alias-tuples; every group must have one alias present."""
    return all(any(a.strip().lower() in hm for a in group) for group in required)


def read_table(path, *required):
    """
    Read a tabular export into (rows, header_map) with 0-based column indices.

    Accepts **.csv or .xlsx/.xlsm for any of the three inputs.** That matters: the
    source map has the employee master coming off the ERP as
    "Employees > Detailed Employee View > Download xlsx" while earlier copies arrived as
    CSV, so either can turn up week to week and neither should break the build.

    `required` is a tuple of alias-groups; every group must be satisfied by the header
    row. For a workbook this also picks the sheet: the first sheet whose header row
    satisfies them (preferring names listed in SHEET_ALIASES). If nothing matches, it
    fails loudly listing the sheets and headers it saw.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        with open(path, newline='', encoding='utf-8-sig') as f:
            rdr = csv.reader(f)
            try:
                header = next(rdr)
            except StopIteration:
                raise SystemExit('ERROR: %s is empty.' % os.path.basename(path))
            hm = _header_map(header)
            rows = [r for r in rdr if any(str(c).strip() for c in r)]
        if required and not _satisfies(hm, required):
            raise SystemExit(
                'ERROR: %s is missing a required column.\n'
                '       Needed one of each: %s\n'
                '       Headers present: %s'
                % (os.path.basename(path), [list(g) for g in required], sorted(hm)))
        return rows, hm

    if ext not in ('.xlsx', '.xlsm'):
        raise SystemExit('ERROR: unsupported file type: %s' % path)

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        preferred = [n for alias in SHEET_ALIASES.values() for n in alias]
        order = sorted(wb.sheetnames, key=lambda n: (n not in preferred,))
        seen = {}
        for name in order:
            ws = wb[name]
            it = ws.iter_rows(values_only=True)
            try:
                header = next(it)
            except StopIteration:
                continue
            hm = _header_map(header)
            seen[name] = sorted(hm)
            if required and not _satisfies(hm, required):
                continue
            rows = [list(r) for r in it if any(c is not None and str(c).strip() != '' for c in r)]
            return rows, hm
        raise SystemExit(
            'ERROR: no sheet in %s has the required columns.\n'
            '       Needed one of each: %s\n'
            '       Sheets and their headers: %s'
            % (os.path.basename(path), [list(g) for g in required], seen))
    finally:
        wb.close()


def hc(hm, *names, tab='', required=True):
    """Resolve a column index by header name, trying each alias in order."""
    for n in names:
        k = hm.get(n.strip().lower())
        if k is not None:
            return k
    if required:
        raise SystemExit(
            "ERROR: [%s] required column not found; tried %s.\n"
            "       Headers present: %s\n"
            "       Add the new header name to the hc() call in build_dashboards.py."
            % (tab, list(names), sorted(hm)))
    return None


def cell(row, idx):
    """Value at column index `idx` of a row, tolerating short rows and blanks."""
    if idx is None or row is None or idx >= len(row):
        return None
    v = row[idx]
    if v is None:
        return None
    return v if str(v).strip() != '' else None


# ------------------------------------------------------------------ file sniffing
# What each export is recognised by. Checked in this order, and deliberately on
# CONTENT not filename or extension -- the ERP stamps a drifting number into every
# name ("Timesheets 9.csv", "Employees 32 1.csv", "Staff expenses 22.xlsx") and the
# same view can be downloaded as either CSV or xlsx.
_SIGNATURES = (
    ('expenses',   (('ER #', 'ER#'),)),
    ('timesheets', (('GL post date', 'GL Posting Date', 'GL Post Date'), ('Begin date', 'Begin Date', 'Begin'))),
    ('employees',  (('Supervisor name', 'Supervisor'),)),
)

# Required columns per input, also used to pick the right sheet inside a workbook.
REQUIRED = {
    'expenses':   (('Employee ID',), ('ER #', 'ER#'), ('State',)),
    'timesheets': (('Employee ID',), ('GL post date', 'GL Posting Date', 'GL Post Date'), ('State',)),
    'employees':  (('Employee ID',), ('Supervisor name', 'Supervisor')),
}


def _classify(path):
    """Which of the three exports is this? Reads only the header row(s)."""
    ext = os.path.splitext(path)[1].lower()
    maps = []
    if ext == '.csv':
        with open(path, newline='', encoding='utf-8-sig') as f:
            try:
                maps.append(_header_map(next(csv.reader(f))))
            except StopIteration:
                pass
    elif ext in ('.xlsx', '.xlsm'):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for name in wb.sheetnames:
                try:
                    maps.append(_header_map(next(wb[name].iter_rows(values_only=True))))
                except StopIteration:
                    continue
        finally:
            wb.close()
    else:
        raise SystemExit('ERROR: unsupported file type: %s' % path)
    for kind, sig in _SIGNATURES:
        for hm in maps:
            if _satisfies(hm, sig):
                return kind, sorted(set().union(*maps)) if maps else []
    return None, sorted(set().union(*maps)) if maps else []


def resolve_sources(*paths):
    """
    Work out which supplied path is which, by content rather than filename or
    extension. Returns (timesheets, employees, expenses) paths.
    """
    found = {'timesheets': None, 'employees': None, 'expenses': None}
    for p in paths:
        if p is None:
            continue
        if not os.path.exists(p):
            raise SystemExit('ERROR: no such file: %s' % p)
        kind, headers = _classify(p)
        if kind is None:
            raise SystemExit(
                'ERROR: cannot tell what %s is. Expected the expense export (has "ER #"), '
                'the timesheet export (has "GL post date" + "Begin date"), or the employee '
                'master (has "Supervisor name").\n       Headers found: %s'
                % (os.path.basename(p), headers))
        if found[kind]:
            raise SystemExit('ERROR: two files look like the %s export: %s and %s'
                             % (kind, os.path.basename(found[kind]), os.path.basename(p)))
        found[kind] = p
    missing = [k for k, v in found.items() if not v]
    if missing:
        raise SystemExit(
            'ERROR: missing the %s file. Pass all three, in any order (.csv or .xlsx both fine):\n'
            '       python build_dashboard_html.py "Timesheets <n>.csv" '
            '"Employees <n>.csv" "Staff expenses <n>.xlsx"'
            % ' and '.join(missing))
    return found['timesheets'], found['employees'], found['expenses']


# ------------------------------------------------------------------ age gate
def _today():
    """Today's date, or the FPA_AS_OF=YYYY-MM-DD override for reproducible rebuilds."""
    v = os.environ.get('FPA_AS_OF')
    return datetime.datetime.strptime(v, '%Y-%m-%d').date() if v else datetime.date.today()


def latest_activity(rows, *keys):
    """Latest real (non-future) activity date across `keys` -- the export date."""
    today = _today()
    ds = []
    for x in rows:
        for k in keys:
            v = x.get(k)
            if isinstance(v, datetime.datetime):
                if v.date() <= today:
                    ds.append(v.date())
                break
    return max(ds) if ds else None


def outstanding_cutoff(as_of):
    """
    First date that is NOT old enough to show. A row dated strictly before this has
    been outstanding more than MIN_DAYS_OUTSTANDING days.
    """
    if not DAYS_OUTSTANDING_ONLY or as_of is None:
        return None
    return as_of - datetime.timedelta(days=MIN_DAYS_OUTSTANDING)


def days_outstanding(x, as_of, keys):
    """Days the row has been outstanding, on the first of `keys` it carries."""
    if as_of is None:
        return None
    for k in keys:
        v = x.get(k)
        if isinstance(v, datetime.datetime):
            return (as_of - v.date()).days
    return None


def _is_outstanding(x, as_of, keys):
    """
    True if the row has been outstanding long enough to show. Rows carrying no usable
    date are KEPT -- they are open and undateable, and hiding them would lose a real
    outstanding item.
    """
    if not DAYS_OUTSTANDING_ONLY or as_of is None:
        return True
    d = days_outstanding(x, as_of, keys)
    return True if d is None else d > MIN_DAYS_OUTSTANDING


# ------------------------------------------------------------------ EMPLOYEES
def build_employees(path):
    """
    The employee master: Employee ID -> supervisor / location / legacy / flags.
    Replaces the old 'Employee (ERP)' and 'ERP Employee Data' tabs.
    """
    rows, hm = read_table(path, *REQUIRED['employees'])
    tab = os.path.basename(path)
    I_ID   = hc(hm, 'Employee ID', tab=tab)
    I_SUP  = hc(hm, 'Supervisor name', 'Supervisor', tab=tab)
    I_LOC  = hc(hm, 'Location ID', tab=tab)
    I_LEG  = hc(hm, 'Legacy ID', tab=tab, required=False)
    I_NE   = hc(hm, 'Not Employed', tab=tab, required=False)
    I_TTL  = hc(hm, 'Title', tab=tab, required=False)
    I_DEPT = hc(hm, 'Department ID', tab=tab, required=False)
    I_NAME = hc(hm, 'Primary contact name', tab=tab, required=False)

    emps = {}
    for r in rows:
        emp = cell(r, I_ID)
        if emp is None:
            continue
        loc_id = cell(r, I_LOC)
        loc_id = str(loc_id).strip() if loc_id is not None else None
        ne = cell(r, I_NE)
        emps[str(emp).strip()] = dict(
            supervisor=clean_name(cell(r, I_SUP)),
            name=clean_name(cell(r, I_NAME)),
            legacy=cell(r, I_LEG),
            location_id=loc_id,
            location=LOCATION_NAMES.get(loc_id, loc_id),
            not_employed=(str(ne).strip().lower() == 'true') if ne is not None else None,
            title=cell(r, I_TTL),
            department=cell(r, I_DEPT),
        )
    if not emps:
        raise SystemExit('ERROR: employee master %s produced no rows.' % tab)
    return emps


# ------------------------------------------------------------------ TIMESHEETS
def build_timesheets(path, emps, as_of=None):
    rows_in, hm = read_table(path, *REQUIRED['timesheets'])
    tab = os.path.basename(path)
    I_ID    = hc(hm, 'Employee ID', tab=tab)
    I_NAME  = hc(hm, 'Employee name', 'Employee Name', tab=tab)
    I_BEGIN = hc(hm, 'Begin date', 'Begin Date', 'Begin', tab=tab)
    I_END   = hc(hm, 'End date', 'End Date', 'End', tab=tab)
    I_GL    = hc(hm, 'GL post date', 'GL Post Date', 'GL Posting Date', tab=tab)
    I_STATE = hc(hm, 'State', tab=tab)

    rows = []
    for r in rows_in:
        emp = cell(r, I_ID)
        if emp is None:
            continue
        emp = str(emp).strip()
        e = emps.get(emp, {})
        rows.append(dict(
            emp_id=emp,
            emp_name=clean_name(cell(r, I_NAME)) or e.get('name'),
            supervisor=e.get('supervisor'),
            legacy=e.get('legacy'),
            location=e.get('location'),
            location_id=e.get('location_id'),
            not_employed=e.get('not_employed'),
            title=e.get('title'),
            begin=as_date(cell(r, I_BEGIN)),
            end=as_date(cell(r, I_END)),
            gl=as_date(cell(r, I_GL)),
            state=cell(r, I_STATE),
            in_master=emp in emps,
        ))

    as_of = as_of or AS_OF or latest_activity(rows, 'gl', 'end')
    cutoff = outstanding_cutoff(as_of)

    def passes(x):
        # NOTE: no Target != 0 test any more -- the new sources carry no target.
        # Non-billable/corporate people are excluded by name in build_dashboard_html.py.
        if x['state'] not in OPEN_STATES:
            return False
        return _is_outstanding(x, as_of, AGE_KEYS_TS)

    kept = [x for x in rows if passes(x)]
    pivot = defaultdict(lambda: defaultdict(int))
    for x in kept:
        pivot[x['supervisor']][x['state']] += 1
    return dict(rows=rows, kept=kept, pivot=pivot, as_of=as_of, cutoff=cutoff)


# -------------------------------------------------------------------- EXPENSES
def build_expenses(path, emps, as_of=None):
    rows_in, hm = read_table(path, *REQUIRED['expenses'])
    tab = os.path.basename(path)
    C_ID     = hc(hm, 'Employee ID', tab=tab)
    C_NAME   = hc(hm, 'Employee Name', 'Employee name', tab=tab)
    C_ER     = hc(hm, 'ER #', 'ER#', tab=tab)
    C_REASON = hc(hm, 'Reason for expense', 'Reason', tab=tab)
    C_STATE  = hc(hm, 'State', tab=tab)
    C_FILED  = hc(hm, 'Date Filed', tab=tab)
    C_GL     = hc(hm, 'GL Posting Date', 'GL Post Date', tab=tab)
    C_AMT    = hc(hm, 'Amount', tab=tab)
    C_LM     = hc(hm, 'Last Modified', 'Last modified', tab=tab, required=False)
    C_SUMM   = hc(hm, 'Summary State', tab=tab, required=False)
    C_DUE    = hc(hm, 'Amount Due', tab=tab, required=False)

    rows = []
    for row in rows_in:
        get = lambda c: cell(row, c)
        er = get(C_ER)
        if er is None or str(er).strip() == '':
            continue
        emp = str(get(C_ID)).strip() if get(C_ID) is not None else ''
        e = emps.get(emp, {})
        rows.append(dict(
            emp_id=emp,
            emp_name=clean_name(get(C_NAME)) or e.get('name'),
            er=str(er).strip(),
            reason=get(C_REASON),
            state=get(C_STATE),
            date_filed=as_date(get(C_FILED)),
            gl=as_date(get(C_GL)),
            last_modified=as_date(get(C_LM)),
            amount=as_num(get(C_AMT)) or 0.0,
            amount_due=as_num(get(C_DUE)),
            summary_state=get(C_SUMM),
            supervisor=e.get('supervisor'),
            location=e.get('location'),
            location_id=e.get('location_id'),
            not_employed=e.get('not_employed'),
            in_master=emp in emps,
        ))

    # Last Modified is the only true timestamp in either file, which makes it the
    # export's clock. build_dashboard_html.py feeds this as_of into build_timesheets.
    as_of = as_of or AS_OF or latest_activity(rows, 'last_modified', 'gl', 'date_filed')
    cutoff = outstanding_cutoff(as_of)

    def passes(x):
        if x['state'] not in OPEN_STATES:
            return False
        return _is_outstanding(x, as_of, AGE_KEYS_EX)

    kept = [x for x in rows if passes(x)]
    pivot = defaultdict(lambda: defaultdict(int))
    for x in kept:
        pivot[x['supervisor']][x['state']] += 1
    return dict(rows=rows, kept=kept, pivot=pivot, as_of=as_of, cutoff=cutoff)


# ------------------------------------------------------------------ self-check
def show(name, res, exp_draft=None, exp_sub=None):
    p = res['pivot']
    d = sum(v.get('Draft', 0) for v in p.values())
    s = sum(v.get('Submitted', 0) for v in p.values())
    n_open = sum(1 for x in res['rows'] if x.get('state') in OPEN_STATES)
    unresolved = sorted({x['emp_id'] for x in res['kept'] if not x.get('in_master')})
    print('\n===== %s =====' % name)
    print('as of: %s  | show items dated before: %s  (outstanding > %d days)'
          % (res.get('as_of'), res.get('cutoff'), MIN_DAYS_OUTSTANDING))
    print('total rows read: %d  open Draft/Submitted: %d  kept as outstanding: %d'
          % (len(res['rows']), n_open, len(res['kept'])))
    print('Draft=%d  Submitted=%d  Total=%d  supervisors=%d  people=%d'
          % (d, s, d + s, len(p), len({str(x['emp_id']) for x in res['kept']})))
    if unresolved:
        print('WARNING: %d kept row(s) have an Employee ID missing from the employee '
              'master (no supervisor/PMO): %s' % (len(unresolved), ', '.join(unresolved[:10])))
    if exp_draft is not None:
        print('   expected %d/%d  %s' % (exp_draft, exp_sub,
              '\u2713' if (d == exp_draft and s == exp_sub) else '\u2717 MISMATCH'))
    rows = sorted(p.items(), key=lambda kv: -(kv[1].get('Draft', 0) + kv[1].get('Submitted', 0)))
    print('Supervisor'.ljust(26), 'Draft  Sub  Total')
    for sup, v in rows:
        dd, ss = v.get('Draft', 0), v.get('Submitted', 0)
        print(str(sup).ljust(26), '%5d%5d%6d' % (dd, ss, dd + ss))


if __name__ == '__main__':
    import sys
    if len(sys.argv) < 4:
        raise SystemExit(
            'usage: python build_dashboards.py <Timesheets.csv> <Employees.csv> '
            '<Staff expenses.xlsx>   (any order)\n'
            "       This week's exports live in:\n"
            '       demo_data/Weekly and Monthly Updates/<M.D.YYYY>/'
            'Timesheets Draft & Submitted/')
    ts_file, emp_file, ex_file = resolve_sources(*sys.argv[1:4])
    emps = build_employees(emp_file)
    print('employee master: %d people from %s' % (len(emps), os.path.basename(emp_file)))
    # Expenses first: its Last Modified is the clock for both files.
    ex = build_expenses(ex_file, emps)
    ts = build_timesheets(ts_file, emps, as_of=ex.get('as_of'))
    show('TIMESHEETS  (Draft/Submitted, GL post date more than %d days outstanding)'
         % MIN_DAYS_OUTSTANDING, ts)
    show('EXPENSE REPORTS  (Draft/Submitted, GL Posting Date more than %d days outstanding)'
         % MIN_DAYS_OUTSTANDING, ex)
