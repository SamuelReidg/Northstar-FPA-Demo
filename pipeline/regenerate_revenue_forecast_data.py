#!/usr/bin/env python3
"""
regenerate_revenue_forecast_data.py
Generates revenue_forecast_data.js (window.REVENUE_FORECAST) for the Revenue
Forecast dashboard from the monthly "Revenue Forecast.xlsx" export.
Source sheet: "OTC-Project Profitability CM to" (headers rows 2-5, data row 7+).
All figures are ACTUALS; the forward forecast is entered by PMs in the dashboard.
Row schema: [id,name,pm,location,region,status,billing,contract,recITD,remaining,recMTD,recPrev,cosITD,suggestRate]
Usage: python regenerate_revenue_forecast_data.py <xlsx> [out.js] [entries.csv]

Optional entries.csv: PM forecast entries in the dashboard's own CSV Export format
(Project ID, Name, PM, Location, Avg Rate, 3 x [Hrs, Cost, Exp], Exp Markup %). When given,
they are embedded as `seed` and become the dashboard's starting values for any project the
viewer has not edited in this browser (their own localStorage entries always win).

Reporting date: read from the export's period-date header (row 6; row 4 checked too), else
today — or FPA_AS_OF=YYYY-MM-DD when set, for reproducible builds.
"""
import sys, os, csv, json, datetime
import openpyxl

SHEET_HINT = "OTC-Project Profitability"
INTL = {"International","NH Gulf (USD)","NH Gulf","Global LLC International","Island"}
# Rows of the export that are not projects: header caption, group header, reserve and totals
NON_PROJECT_IDS = {"Project ID","All Projects","WIP Reserve Project","Total All Projects","All Locations"}
# Statuses to drop entirely (dead / inactive); everything else is kept
EXCLUDE_STATUS = {"Canceled","Completed","On Hold","Legacy","Legacy-Hold","Cost Tracking on Closed Project","Final Billing/Invoice","Work Hold"}

def region_of(loc):
    if not loc: return "Domestic"
    l = str(loc).strip()
    if "Meridian" in l: return "Meridian"
    if l in INTL: return "International"
    return "Domestic"

def num(v):
    try:
        if v is None or v == "": return 0.0
        return float(v)
    except (TypeError, ValueError): return 0.0

def today():
    ov = os.environ.get("FPA_AS_OF")
    if ov:
        return datetime.datetime.strptime(ov, "%Y-%m-%d").date()
    return datetime.date.today()

def header_date(ws, col):
    for r in (6, 4):
        v = ws.cell(r, col).value
        if isinstance(v, datetime.datetime): return v
    return None

def load_seed(path, ids):
    """PM forecast entries (dashboard CSV Export format) -> {id: {rate, mk, m:[{h,c,x} x3]}}."""
    seed = {}
    if not path: return seed
    with open(path, newline="", encoding="utf-8-sig") as f:
        rd = csv.reader(f); next(rd, None)
        for c in rd:
            if len(c) < 14 or c[0] not in ids: continue
            n = [num(v) for v in c[4:15]] + [0.0] * 11
            seed[c[0]] = {"rate": n[0], "mk": n[10],
                          "m": [{"h": n[1+3*k], "c": n[2+3*k], "x": n[3+3*k]} for k in range(3)]}
    return seed

def main():
    if len(sys.argv) < 2:
        print("usage: python regenerate_revenue_forecast_data.py <xlsx> [out.js] [entries.csv]"); sys.exit(1)
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "revenue_forecast_data.js"
    entries = sys.argv[3] if len(sys.argv) > 3 else None
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = next((s for s in wb.worksheets if SHEET_HINT.lower() in s.title.lower()), wb.active)

    as_of = header_date(ws, 8)
    as_of_str = as_of.strftime("%B %d, %Y") if isinstance(as_of, datetime.datetime) else today().strftime("%B %d, %Y")
    prev = header_date(ws, 17)
    prev_str = prev.strftime("%B %Y") if isinstance(prev, datetime.datetime) else ""

    rows = []; tot_contract = tot_recITD = tot_cos = 0.0
    for r in range(7, ws.max_row + 1):
        pid = ws.cell(r, 1).value
        if pid in (None, "") or str(pid).strip() in NON_PROJECT_IDS: continue
        name = ws.cell(r, 2).value or ""
        billing = ws.cell(r, 3).value or "—"
        status = ws.cell(r, 4).value or "—"
        if str(status).strip() in EXCLUDE_STATUS:
            continue
        _idn = (str(pid) + " " + str(ws.cell(r, 2).value or "")).upper()
        if "OVHD" in _idn or "OVERHEAD" in _idn:
            continue
        contract = num(ws.cell(r, 8).value)
        recMTD = num(ws.cell(r, 11).value); hrsMTD = num(ws.cell(r, 14).value)
        recPrev = num(ws.cell(r, 17).value); hrsPrev = num(ws.cell(r, 20).value)
        recITD = num(ws.cell(r, 23).value); cosITD = num(ws.cell(r, 24).value)
        loc = ws.cell(r, 27).value or "—"; pm = ws.cell(r, 29).value or "—"
        rate = recMTD/hrsMTD if hrsMTD > 0 else (recPrev/hrsPrev if hrsPrev > 0 else 0.0)
        rate = round(rate / 5.0) * 5
        rows.append([str(pid), str(name).strip(), str(pm).strip(), str(loc).strip(),
            region_of(loc), str(status).strip(), str(billing).strip(),
            round(contract), round(recITD), round(contract - recITD),
            round(recMTD), round(recPrev), round(cosITD), int(rate)])
        tot_contract += contract; tot_recITD += recITD; tot_cos += cosITD

    meta = {"asOf": as_of_str, "prevMonth": prev_str,
            "months": ["October 2026","November 2026","December 2026"],
            "monthsShort": ["October","November","December"], "n": len(rows)}
    payload = {"meta": meta, "rows": rows}
    seed = load_seed(entries, {x[0] for x in rows})
    if entries:
        meta["seedEntries"] = len(seed)
        payload["seed"] = seed
    js = "window.REVENUE_FORECAST = " + json.dumps(payload, separators=(",", ":")) + ";\n"
    with open(out, "w", encoding="utf-8") as f: f.write(js)

    inprog = [x for x in rows if x[5] == "In Progress"]
    ip_c = sum(x[7] for x in inprog); ip_r = sum(x[8] for x in inprog)
    print("wrote %s  (%d projects, %.0f KB)" % (out, len(rows), len(js)/1024))
    print("ALL         contract=$%.1fM recITD=$%.1fM cos=$%.1fM" % (tot_contract/1e6, tot_recITD/1e6, tot_cos/1e6))
    print("In Progress n=%d contract=$%.1fM recITD=$%.1fM remaining=$%.1fM" % (len(inprog), ip_c/1e6, ip_r/1e6, (ip_c-ip_r)/1e6))
    print("as of:", as_of_str, "| prev:", prev_str)
    if entries:
        frev = [sum(e["m"][k]["h"] * e["rate"] + e["m"][k]["x"] * (1 + e["mk"] / 100.0) for e in seed.values()) for k in range(3)]
        print("seeded entries: %d  3-mo expected revenue=$%.2fM (%s)" % (len(seed), sum(frev)/1e6,
              " / ".join("$%.2fM" % (v/1e6) for v in frev)))
        over = [i for i, e in seed.items() if sum(e["m"][k]["h"] * e["rate"] + e["m"][k]["x"] * (1 + e["mk"]/100.0) for k in range(3))
                > max(0, next(x[9] for x in rows if x[0] == i)) + 1]
        print("  entries forecasting past remaining contract value: %d" % len(over))

if __name__ == "__main__":
    main()
